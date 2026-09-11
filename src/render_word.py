"""Offline deterministic Markdown-to-DOCX renderer."""
from __future__ import annotations
import json, os, re, tempfile, zipfile
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from src.markdown_images import image_blocks, image_syntax_errors, safe_asset_path
from src.delivery_text import formal_markdown, forbidden_delivery_tokens
HEADING=re.compile(r"^(#{1,6})\s+(.+)$"); NUM=re.compile(r"\b\d+(?:\.\d+)?\s*(?:MWp|MWh|MW|kV|kW|V|W|A|mm²|mm2|%)\b",re.I)
class RenderWordError(RuntimeError): pass
def font(run,name='宋体',size=12,bold=False,italic=False):
 run.font.name=name;run.font.size=Pt(size);run.bold=bold;run.italic=italic;r=run._element.get_or_add_rPr();f=r.rFonts or OxmlElement('w:rFonts');f.set(qn('w:eastAsia'),name);f.set(qn('w:ascii'),'Arial');r.insert(0,f)
def inline(p,s):
 for x in re.split(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)",s):
  b=x.startswith('**'); i=not b and x.startswith('*'); code=x.startswith('`'); font(p.add_run(x[2:-2] if b else x[1:-1] if i or code else x),'Consolas' if code else '宋体',10 if code else 12,b,i)
def sources(md, rejected=()):
 text=md.read_text(encoding='utf8'); errors=image_syntax_errors(text)
 if errors: raise RenderWordError('invalid proposal image syntax: ' + '；'.join(errors))
 refs=[]
 for block in image_blocks(text):
  if block.asset_path in rejected: continue
  p=safe_asset_path(md,block.asset_path)
  if not p: raise RenderWordError(f'line {block.line_number}: missing or unsafe image: {block.asset_path}')
  refs.append((block.caption,p))
 return refs,list(dict.fromkeys(p for _,p in refs))
def render_word(input_path,output_path,*,formal=False,image_decisions_path=None):
 md=Path(input_path).resolve();out=Path(output_path).resolve()
 if not md.is_file() or md.suffix.lower()!='.md':raise RenderWordError('--input must be an existing Markdown file')
 text=md.read_text(encoding='utf8'); rejected=set()
 if image_decisions_path:
  rejected={p for p,v in json.loads(Path(image_decisions_path).read_text(encoding='utf8')).get('images',{}).items() if not v.get('use_in_word',v.get('status')=='used')}
 if formal: text,_=formal_markdown(text)
 refs,unique=sources(md,rejected);doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Cm(2.5);sec.left_margin=sec.right_margin=Cm(2.7); hs=ts=0;lines=text.splitlines();i=0
 while i<len(lines):
  line=lines[i];h=HEADING.match(line); blocks=image_blocks(line)
  # Consume every recognised Markdown image. Rejected figures must not fall
  # through to ordinary text and leak their raw Markdown into the DOCX.
  if blocks:
   block=blocks[0]
   if block.asset_path not in rejected:
    alt,path=block.caption,safe_asset_path(md,block.asset_path);p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(path),width=Cm(15));
    if alt:q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;font(q.add_run(alt),size=10)
  elif h:
   lv=len(h.group(1));p=doc.add_heading(level=lv);p.paragraph_format.keep_with_next=True;font(p.add_run(h.group(2)),'黑体',22 if lv==1 else 18 if lv==2 else 15,True);hs+=1
  elif line.strip().startswith('|') and i+1<len(lines) and '-' in lines[i+1]:
   block=[line];i+=1
   while i+1<len(lines) and lines[i+1].strip().startswith('|'):i+=1;block.append(lines[i])
   rows=[[x.strip() for x in z.strip().strip('|').split('|')] for z in block if '-' not in z];t=doc.add_table(rows=len(rows),cols=len(rows[0]));t.style='Table Grid'
   for r,row in enumerate(rows):
    for c,v in enumerate(row):inline(t.cell(r,c).paragraphs[0],v)
   ts+=1
  elif re.match(r'^\s*(?:[-*+] |\d+\. )',line):p=doc.add_paragraph(style='List Number' if re.match(r'^\s*\d+\.',line) else 'List Bullet');inline(p,re.sub(r'^\s*(?:[-*+] |\d+\. )','',line))
  elif line.strip():p=doc.add_paragraph();inline(p,line.lstrip('> ').strip())
  i+=1
 out.parent.mkdir(parents=True,exist_ok=True)
 with tempfile.TemporaryDirectory(dir=out.parent,prefix='render-word-') as d:
  temp=Path(d)/out.name;doc.save(temp)
  with zipfile.ZipFile(temp) as z:media=[n for n in z.namelist() if n.startswith('word/media/')];assert '[Content_Types].xml' in z.namelist() and 'word/document.xml' in z.namelist()
  check=Document(temp);body='\n'.join(p.text for p in check.paragraphs)+'\n'+'\n'.join(c.text for t in check.tables for r in t.rows for c in r.cells)
  if any(v not in body for _,v in HEADING.findall(text)) or len(check.tables)!=ts or len(check.inline_shapes)!=len(refs):raise RenderWordError('local DOCX validation failed')
  counts={'markdown_image_path_count':len(refs),'standalone_image_block_count':len(refs),'embedded_image_count':len(check.inline_shapes),'unique_asset_count':len(unique)}
  if len(set(counts.values())) != 1: raise RenderWordError('image count mismatch: '+json.dumps(counts))
  forbidden=forbidden_delivery_tokens(body) if formal else []
  if forbidden: raise RenderWordError('formal delivery tokens remain: '+','.join(forbidden))
  report={'status':'PASS','backend':'local_python_docx','formal_delivery':formal,'input_markdown':md.name,'output_docx':out.name,'source_heading_count':sum(bool(HEADING.match(x)) for x in lines),'rendered_heading_count':hs,'source_table_count':ts,'rendered_table_count':len(check.tables),'source_image_reference_count':len(refs),'unique_source_image_count':len(unique),**counts,'docx_media_count':len(media),'inline_shape_count':len(check.inline_shapes),'checked_numeric_tokens':NUM.findall(text),'warnings':[],'errors':[]};os.replace(temp,out);out.with_suffix('.render_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
 return out.with_suffix('.render_report.json')
