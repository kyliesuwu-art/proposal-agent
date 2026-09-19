from __future__ import annotations
import re, zipfile
from pathlib import Path
from docx import Document
from .common import check, digest, text_counts, zip_is_safe
from .models import ArtifactReport, EvaluationStatus, Severity, aggregate
from .profiles import Profile

def evaluate_docx(value: str | Path, profile: Profile) -> ArtifactReport:
    path=Path(value); report=ArtifactReport("docx",str(path),None,None,profile.name)
    if not path.is_file(): report.checks.append(check("DOCX_FILE_MISSING",EvaluationStatus.FAIL,Severity.ERROR,"docx","DOCX file is missing",location=str(path)));report.status=EvaluationStatus.FAIL;return report
    report.file_size,report.sha256=path.stat().st_size,digest(path); ok,detail=zip_is_safe(path,"word/document.xml")
    if not ok: report.checks.append(check("DOCX_OOXML_INVALID",EvaluationStatus.FAIL,Severity.ERROR,"docx","Invalid DOCX package",evidence=detail));report.status=EvaluationStatus.FAIL;return report
    try: doc=Document(path)
    except Exception as exc: report.checks.append(check("DOCX_PARSE_FAILED",EvaluationStatus.FAIL,Severity.ERROR,"docx","python-docx cannot open file",evidence=type(exc).__name__));report.status=EvaluationStatus.FAIL;return report
    text="\n".join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells]); internal,source,paths=text_counts(text)
    headings=[p for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")]; media=[]
    with zipfile.ZipFile(path) as z: media=[n for n in z.namelist() if n.startswith("word/media/")]
    section=doc.sections[0] if doc.sections else None; page_size=None
    if section: page_size={"width_mm":round(section.page_width.mm,1),"height_mm":round(section.page_height.mm,1)}
    report.metrics={"paragraph_count":len(doc.paragraphs),"nonempty_paragraph_count":sum(bool(p.text.strip()) for p in doc.paragraphs),"heading_count":len(headings),"table_count":len(doc.tables),"image_count":len(media),"internal_label_count":internal,"markdown_residue_count":len(re.findall(r"(?:^|\s)(?:#{1,6}\s|\*\*|`|!\[)",text)),"source_leak_count":source+paths,"page_size":page_size,"visual_checks_status":"NOT_RUN"}
    if not text.strip(): report.checks.append(check("DOCX_NO_BODY",EvaluationStatus.FAIL,Severity.ERROR,"docx","DOCX has no body text"))
    if profile.delivery and internal: report.checks.append(check("DOCX_INTERNAL_LABEL",EvaluationStatus.FAIL,Severity.ERROR,"docx","internal labels found",metric_value=internal))
    if profile.delivery and source+paths: report.checks.append(check("DOCX_SOURCE_LEAK",EvaluationStatus.FAIL,Severity.ERROR,"docx","source chain or local path found",metric_value=source+paths))
    if report.metrics["markdown_residue_count"]: report.checks.append(check("DOCX_MARKDOWN_RESIDUE",EvaluationStatus.PASS_WITH_WARNINGS,Severity.WARNING,"docx","Markdown residue found",metric_value=report.metrics["markdown_residue_count"]))
    report.checks.extend([check("DOCX_PARSED",EvaluationStatus.PASS,Severity.INFO,"docx","DOCX parsed"),check("DOCX_VISUAL_RENDER",EvaluationStatus.NOT_RUN,Severity.INFO,"docx","visual layout requires rendered pages")]);report.status=aggregate(report.checks);return report
