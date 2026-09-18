from __future__ import annotations
import hashlib, re
from collections import Counter
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from .common import check,digest,zip_is_safe
from .models import ArtifactReport,EvaluationStatus,Severity,aggregate
from .profiles import Profile

EMU=914400
def evaluate_pptx(value: str|Path,profile:Profile)->ArtifactReport:
 path=Path(value);report=ArtifactReport("pptx",str(path),None,None,profile.name)
 if not path.is_file(): report.checks.append(check("PPTX_FILE_MISSING",EvaluationStatus.FAIL,Severity.ERROR,"pptx","PPTX file is missing",location=str(path)));report.status=EvaluationStatus.FAIL;return report
 report.file_size,report.sha256=path.stat().st_size,digest(path);ok,detail=zip_is_safe(path,"ppt/presentation.xml")
 if not ok: report.checks.append(check("PPTX_OOXML_INVALID",EvaluationStatus.FAIL,Severity.ERROR,"pptx","Invalid PPTX package",evidence=detail));report.status=EvaluationStatus.FAIL;return report
 try: prs=Presentation(path)
 except Exception as exc: report.checks.append(check("PPTX_PARSE_FAILED",EvaluationStatus.FAIL,Severity.ERROR,"pptx","python-pptx cannot open file",evidence=type(exc).__name__));report.status=EvaluationStatus.FAIL;return report
 out=low=blank=oob=small=0; signatures=[]; page_errors=0; all_images=[]
 for number,slide in enumerate(prs.slides,1):
  texts=[]; shapes=[]; has_content=False
  for shape in slide.shapes:
   if shape.left<0 or shape.top<0 or shape.left+shape.width>prs.slide_width or shape.top+shape.height>prs.slide_height:oob+=1
   kind=str(shape.shape_type); shapes.append((kind,round(shape.left/prs.slide_width,2),round(shape.top/prs.slide_height,2),round(shape.width/prs.slide_width,2),round(shape.height/prs.slide_height,2)))
   if getattr(shape,"has_text_frame",False):
    texts.append(shape.text or ""); has_content|=bool(shape.text.strip())
    for paragraph in shape.text_frame.paragraphs:
     for run in paragraph.runs:
      if run.font.size and run.font.size.pt<profile.min_font_pt:small+=1
   if shape.shape_type==MSO_SHAPE_TYPE.PICTURE:
    has_content=True; out+=1
    try:
     blob=shape.image.blob; key=hashlib.sha256(blob).hexdigest(); all_images.append(key)
     pxw,pxh=shape.image.size; ppi=min(pxw/(shape.width/EMU),pxh/(shape.height/EMU))
     if ppi<profile.min_ppi_error: low+=1; report.checks.append(check("PPTX_IMAGE_PPI_ERROR",EvaluationStatus.PASS_WITH_WARNINGS,Severity.WARNING,"pptx",f"slide {number} image PPI is low",location=f"slide {number}",metric_value=round(ppi,1),expected_value=profile.min_ppi_error))
     elif ppi<profile.min_ppi_warning: low+=1
    except Exception: pass
  if not has_content: blank+=1
  signatures.append(tuple(sorted(shapes)))
  if number>1 and str(number) not in " ".join(texts): page_errors+=1
 counts=Counter(signatures); repeated=sum(1 for count in counts.values() if count>1)
 report.metrics={"slide_count":len(prs.slides),"image_count":out,"low_resolution_image_count":low,"out_of_bounds_shape_count":oob,"missing_logo_slide_count":None,"duplicate_logo_slide_count":None,"missing_footer_slide_count":None,"page_number_error_count":page_errors,"small_font_count":small,"blank_slide_count":blank,"repeated_layout_group_count":repeated,"duplicate_image_count":len(all_images)-len(set(all_images)),"visual_checks_status":"NOT_RUN"}
 if blank:report.checks.append(check("PPTX_BLANK_SLIDE",EvaluationStatus.FAIL,Severity.ERROR,"pptx","blank slide found",metric_value=blank))
 if oob:report.checks.append(check("PPTX_SHAPE_OUT_OF_BOUNDS",EvaluationStatus.FAIL,Severity.ERROR,"pptx","shape exceeds canvas",metric_value=oob))
 if low:report.checks.append(check("PPTX_LOW_RESOLUTION_IMAGE",EvaluationStatus.PASS_WITH_WARNINGS,Severity.WARNING,"pptx","low resolution image risk",metric_value=low))
 if repeated:report.checks.append(check("PPTX_REPEATED_LAYOUT",EvaluationStatus.PASS_WITH_WARNINGS,Severity.WARNING,"pptx","repeated layout group",metric_value=repeated))
 if small:report.checks.append(check("PPTX_SMALL_FONT",EvaluationStatus.PASS_WITH_WARNINGS,Severity.WARNING,"pptx","small font risk",metric_value=small))
 report.checks.extend([check("PPTX_PARSED",EvaluationStatus.PASS,Severity.INFO,"pptx","PPTX parsed"),check("PPTX_VISUAL_RENDER",EvaluationStatus.NOT_RUN,Severity.INFO,"pptx","text overflow and visual collisions require rendered pages")]);report.status=aggregate(report.checks);return report
