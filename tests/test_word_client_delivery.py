import io
import json
from pathlib import Path

from docx import Document

from scripts import run_word_document_director as word_director
from scripts.run_word_document_director import Block, Section, _client_visible_text, _condition_groups, _configure_client_document, _editorial_sections, _explicit_condition_items, _overview_blocks, render_client_delivery


def test_client_delivery_hides_provenance_preserves_standard_and_images(tmp_path: Path):
    from PIL import Image

    source = tmp_path / "internal.docx"
    image = tmp_path / "image.png"
    Image.new("RGB", (40, 20), "navy").save(image)
    doc = Document()
    doc.add_heading("建设方案", level=1)
    doc.add_paragraph("采用 GB/T 1234，容量 500kW。[来源: 内部资料.pdf, 第 2 页]")
    doc.add_paragraph().add_run().add_picture(str(image))
    doc.add_paragraph().add_run().add_picture(str(image))
    doc.add_paragraph("图片来源：内部资料.pdf，第 2 页")
    doc.add_heading("待确认项", level=2)
    doc.add_paragraph("负荷【待确认】")
    doc.add_heading("来源与依据", level=2)
    doc.add_paragraph("内部资料.pdf，第 2 页")
    doc.save(source)
    markdown = "\n".join([
        "医院床位数、总供电容量、实际峰值负荷【待确认】",
        "配电房改造数量、设备选型、实际投资金额【待确认】",
        *[f"事项{index}【待确认】" for index in range(1, 11)],
        "GB/T 1234",
    ])
    output = tmp_path / "client.docx"
    report = render_client_delivery(source, markdown, output)
    body = "\n".join(p.text for p in Document(output).paragraphs)
    assert report["delivery_mode"] == "client"
    assert report["confirmation_item_count"] == 15
    assert "来源与依据" not in body and ".pdf" not in body and "第 2 页" not in body
    assert "GB/T 1234" in body and "实施前需确认事项" in body and "负荷【待确认】" not in body
    assert ".pdf" in "\n".join(p.text for p in Document(source).paragraphs)


def test_client_visible_text_removes_only_bracketed_provenance():
    value = "容量 500kW，适用 GB/T 1234。[来源: 手册.pdf, 第 2 页]"
    cleaned = _client_visible_text(value)
    assert "500kW" in cleaned and "GB/T 1234" in cleaned
    assert "手册.pdf" not in cleaned and "第 2 页" not in cleaned


def test_client_visible_text_removes_only_source_tail():
    value = "容量 500kW，适用 GB/T 1234；来源：手册.pdf，第 2 页"
    cleaned = _client_visible_text(value)
    assert cleaned == "容量 500kW，适用 GB/T 1234"


def test_overview_blocks_extract_existing_text_without_fixed_project_names():
    sections = [
        Section("section_01", "现状", [Block("B1", "paragraph", "现状已经明确。后续不复制。")]),
        Section("section_02", "目标", [Block("B2", "paragraph", "建设目标已经明确。")]),
    ]
    overview = _overview_blocks(sections)
    assert overview[0] == ("项目现状", "现状已经明确。")
    assert "现状；目标" in dict(overview).get("核心建设内容", "")


def test_editorial_transform_moves_confirmation_and_removes_only_internal_label():
    sections = [Section("section_01", "通用章节", [
        Block("B1", "paragraph", "核心目标：保留技术事实。"),
        Block("B2", "paragraph", "参数【待确认】"),
        Block("B3", "paragraph", "无标签正文。"),
    ])]
    transformed, audit = _editorial_sections(sections)
    assert [block.text for block in transformed[0].blocks] == ["保留技术事实。", "参数将在现场勘察和方案深化阶段核实", "无标签正文。"]
    assert any(item["action"] == "derive_condition" for item in audit)
    assert any(item["action"] == "unchanged" for item in audit)


def test_condition_groups_are_generic_and_preserve_every_item():
    values = ["容量与负荷曲线", "现场设备条件", "网络安全要求", "投资测算基础", "项目配合窗口"]
    groups, mapping = _condition_groups(values)
    assert sum(len(items) for _, items in groups) == len(values)
    assert all(item["preserved_key_semantics"] for item in mapping)


def test_editorial_pending_marker_keeps_the_implementation_paragraph_and_derives_condition():
    sections = [Section("section_01", "实施路径", [
        Block("B1", "paragraph", "第二阶段为储能部署与接口预留。具体容量及投资金额【待确认】。"),
    ])]
    transformed, audit = _editorial_sections(sections)
    text = transformed[0].blocks[0].text
    assert "第二阶段为储能部署与接口预留" in text
    assert "待确认" not in text
    assert any(item["action"] == "derive_condition" for item in audit)


def test_editorial_labels_and_provenance_only_captions_do_not_leave_orphan_text():
    sections = [Section("section_01", "通用章节", [
        Block("B1", "paragraph", "【设计建议】建议构建设备状态监测。"),
        Block("B2", "paragraph", "【可选能力】引入巡检机器人。"),
        Block("B3", "paragraph", "架构示意；来源：内部资料.pdf，第1页"),
    ])]
    transformed, _audit = _editorial_sections(sections)
    text = "\n".join(block.text for block in transformed[0].blocks)
    assert "【设计建议】" not in text and "【可选能力】" not in text
    assert "建议构建" in text and "可根据现场条件和运维需求配置巡检机器人" in text
    assert "架构示意" not in text


def test_editorial_conditions_use_dedicated_list_not_implementation_prose_and_a4_is_fixed():
    sections = [
        Section("section_01", "实施路径", [Block("B1", "paragraph", "第二阶段为接口预留。参数【待确认】。")]),
        Section("section_02", "待确认项", [Block("B2", "heading", "运行资料", level=3), Block("B3", "bullet", "负荷曲线【待确认】")]),
    ]
    assert _explicit_condition_items(sections) == [("运行资料", "负荷曲线")]
    document = Document()
    _configure_client_document(document, "通用方案")
    section = document.sections[0]
    assert round(section.page_width.mm) == 210 and round(section.page_height.mm) == 297


def test_editorial_removes_adjacent_duplicate_heading_and_space_before_full_stop():
    sections = [Section("section_01", "配电房数字化改造", [
        Block("B1", "heading", "配电房数字化改造", level=3),
        Block("B2", "heading", "配电房数字化改造", level=3),
        Block("B3", "paragraph", "运行状态达标 。"),
    ])]
    transformed, _audit = _editorial_sections(sections)
    assert [block.text for block in transformed[0].blocks] == ["运行状态达标。"]


def test_editorial_rewrites_malformed_site_confirmation_phrase_generically():
    sections = [Section("section_01", "通用章节", [
        Block("B1", "paragraph", "具体的数据采集点位需结合现场配电拓扑将在现场勘察阶段核实。"),
    ])]
    transformed, _audit = _editorial_sections(sections)
    assert transformed[0].blocks[0].text == "相关数据采集点位将在现场勘察和方案深化阶段核实。"


def test_section_plan_request_disables_thinking_so_json_output_has_a_token_budget(monkeypatch):
    captured = {}

    class Response(io.BytesIO):
        status = 200
        def __enter__(self): return self
        def __exit__(self, *_args): return False

    def urlopen(request, timeout):
        captured["body"] = json.loads(request.data.decode("utf-8"))
        assert timeout == 330
        return Response(b'data: {"choices":[{"delta":{"content":"{\\"ok\\": true}"}}]}\n\ndata: [DONE]\n')

    monkeypatch.setattr(word_director, "env_values", lambda: {"ARK_API_KEY": "test"})
    monkeypatch.setattr(word_director.urllib.request, "urlopen", urlopen)
    result, _call = word_director.model_json("test", "section_01")
    assert result == {"ok": True}
    assert captured["body"]["thinking"] == {"type": "disabled"}
