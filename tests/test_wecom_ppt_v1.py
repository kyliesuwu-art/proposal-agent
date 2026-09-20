from __future__ import annotations

import asyncio
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from pptx import Presentation

from src.wecom.adapter import IncomingMessage, WeComAgentAdapter
from src.wecom.artifact_service import ArtifactService
from src.wecom.control_service import TaskControlService
from src.wecom.models import ArtifactJobStatus, ArtifactType
from src.wecom.ppt_runner import PptResult, ProductionPptRunner
from src.wecom.store import SQLiteTaskStore
from src.wecom.task_service import TaskService
from src.wecom.word_runner import WordResult


class ProposalRunner:
    def run(self, _request: str, output_dir: Path) -> Path:
        output_dir.mkdir(parents=True, exist_ok=True)
        path = output_dir / "proposal.md"
        path.write_text("# Approved\n\n## Confirmations\nOwner【待确认】\n", encoding="utf-8")
        path.write_text("# 已确认方案\n\n## 实施前确认项\n负责人【待确认】\n", encoding="utf-8")
        return path


class WordRunner:
    def run(self, _approved: Path, output_dir: Path, _task_id: str) -> WordResult:
        path = output_dir / "proposal.docx"; path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"PK\x03\x04word")
        return WordResult(path)


class FakePptRunner:
    def __init__(self, status: str = "PASS_WITH_WARNINGS") -> None:
        self.status, self.calls = status, []

    def run(self, approved: Path, task_root: Path, output_dir: Path, task_id: str, job_id: str) -> PptResult:
        self.calls.append((approved, task_root, output_dir, task_id, job_id))
        output_dir.mkdir(parents=True, exist_ok=True)
        deck = output_dir / "proposal.pptx"
        presentation = Presentation(); presentation.slides.add_slide(presentation.slide_layouts[6]); presentation.save(deck)
        return PptResult(deck, evaluation_status=self.status, warnings=("visual_checks_not_run",) if self.status == "PASS_WITH_WARNINGS" else ())


class Transport:
    def __init__(self) -> None: self.replies, self.sent, self.files = [], [], []
    async def reply_text(self, _frame, text): self.replies.append(text)
    async def send_text(self, chatid, text): self.sent.append((chatid, text))
    async def send_file(self, chatid, path): self.files.append((chatid, path))


def harness(tmp_path: Path, runner: FakePptRunner):
    store = SQLiteTaskStore(tmp_path / "wecom.sqlite")
    tasks = TaskService(store, ProposalRunner(), output_root=tmp_path / "tasks")
    artifacts = ArtifactService(store, WordRunner(), runner, task_output_root=tmp_path / "tasks", output_root=tmp_path / "artifacts")
    transport = Transport(); adapter = WeComAgentAdapter(tasks, transport, artifacts, TaskControlService(store, artifacts))
    return store, tasks, artifacts, transport, adapter


def approved(parts, userid="u1", chatid="c1"):
    store, tasks, *_ = parts
    task, _ = tasks.receive(IncomingMessage(userid, chatid, "\u533b\u9662\u4f9b\u914d\u7535\u6539\u9020\u65b9\u6848\uff0c\u91cd\u70b9\u5305\u62ec\u4f9b\u7535\u53ef\u9760\u6027\u548c\u667a\u80fd\u5de1\u68c0"))
    tasks.wait_for_idle(); task, _ = tasks.approve(task.task_id)
    return store.get(task.task_id)


def close(parts):
    store, tasks, artifacts, *_ = parts
    artifacts.shutdown(); tasks.shutdown(); store.close()


def test_ppt_fake_chain_evaluates_before_delivery_and_records_warnings(tmp_path):
    runner = FakePptRunner(); parts = harness(tmp_path, runner)
    try:
        task = approved(parts); store, _tasks, artifacts, transport, adapter = parts
        async def scenario():
            await adapter.handle(IncomingMessage("u1", "c1", "生成PPT", "ppt-1"), object())
            await asyncio.to_thread(artifacts.wait_for_idle); await asyncio.sleep(0)
        asyncio.run(scenario())
        job = store.artifact_jobs_for(task.task_id, ArtifactType.PPTX)[0]
        assert job.status is ArtifactJobStatus.DELIVERED
        assert len(runner.calls) == 1 and len(transport.files) == 1
        warning = Path(job.output_dir) / "evaluation_warnings.json"
        assert json.loads(warning.read_text(encoding="utf-8"))["evaluation_status"] == "PASS_WITH_WARNINGS"
    finally: close(parts)


def test_ppt_evaluation_fail_never_sends_file(tmp_path):
    runner = FakePptRunner("FAIL"); parts = harness(tmp_path, runner)
    try:
        task = approved(parts); store, _tasks, artifacts, transport, adapter = parts
        async def scenario():
            await adapter.handle(IncomingMessage("u1", "c1", "生成PPT", "ppt-fail"), object())
            await asyncio.to_thread(artifacts.wait_for_idle); await asyncio.sleep(0)
        asyncio.run(scenario())
        assert store.artifact_jobs_for(task.task_id, ArtifactType.PPTX)[0].status is ArtifactJobStatus.FAILED
        assert not transport.files
    finally: close(parts)


def test_ppt_resend_does_not_create_job_or_run_model_and_is_isolated(tmp_path):
    runner = FakePptRunner(); parts = harness(tmp_path, runner)
    try:
        task = approved(parts); store, _tasks, artifacts, transport, adapter = parts
        async def scenario():
            await adapter.handle(IncomingMessage("u1", "c1", "生成PPT", "ppt-1"), object())
            await asyncio.to_thread(artifacts.wait_for_idle); await asyncio.sleep(0)
            await adapter.handle(IncomingMessage("u1", "c1", "重发PPT", "ppt-resend"), object())
            await adapter.handle(IncomingMessage("u2", "c1", f"重发PPT {task.task_id}", "other-user"), object())
            await adapter.handle(IncomingMessage("u1", "c2", f"重发PPT {task.task_id}", "other-chat"), object())
        asyncio.run(scenario())
        assert len(store.artifact_jobs_for(task.task_id, ArtifactType.PPTX)) == 1
        assert len(runner.calls) == 1 and len(transport.files) == 2
    finally: close(parts)


def test_ppt_running_job_recovers_as_failed_after_restart(tmp_path):
    runner = FakePptRunner(); parts = harness(tmp_path, runner)
    store, tasks, artifacts, _transport, _adapter = parts
    try:
        task = approved(parts)
        from dataclasses import replace
        now = "2026-09-20T00:00:00+08:00"
        source = Path(task.approved_md_path)
        from src.wecom.models import ArtifactJob
        job = ArtifactJob("runningppt01", task.task_id, ArtifactType.PPTX, ArtifactJobStatus.RUNNING, str(source), "x", str(tmp_path / "artifacts" / task.task_id / "runningppt01"), None, None, None, None, now, now)
        store.create_artifact_job(job); artifacts.shutdown()
        reopened = ArtifactService(store, WordRunner(), runner, task_output_root=tmp_path / "tasks", output_root=tmp_path / "artifacts")
        assert store.get_artifact_job(job.job_id).status is ArtifactJobStatus.FAILED
        reopened.shutdown()
    finally:
        tasks.shutdown(); store.close()


def test_production_runner_uses_disable_model_only_for_offline_rerender(tmp_path, monkeypatch):
    captured = []
    def fake_run(command, **_kwargs):
        captured.append(command)
        output = Path(command[command.index("--output-pptx") + 1]); output.parent.mkdir(parents=True, exist_ok=True)
        with ZipFile(output, "w") as archive:
            archive.writestr("ppt/presentation.xml", "<p:presentation/>")
        report = output.parent / "artifact_evaluation" / "evaluation_report.json"; report.parent.mkdir()
        report.write_text('{"overall_status":"PASS","limitations":[]}', encoding="utf-8")
        return type("Done", (), {"returncode": 0, "stdout": "", "stderr": ""})()
    monkeypatch.setattr("src.wecom.ppt_runner.subprocess.run", fake_run)
    approved = tmp_path / "approved.md"; approved.write_text("# x", encoding="utf-8")
    offline = ProductionPptRunner(tmp_path, existing_scene_dir=tmp_path / "scene")
    offline.run(approved, tmp_path, tmp_path / "offline", "task", "job")
    live = ProductionPptRunner(tmp_path, model_enabled=True, model_scene_command=["approved-scene-producer", "${VISUAL_REFERENCE_ROOT}", "${MODEL_MAX_CALLS}"], model_max_calls=1, visual_reference_root=tmp_path)
    live.run(approved, tmp_path, tmp_path / "live", "task", "job2")
    assert "--disable-model" in captured[0] and "--enable-model" not in captured[0]
    assert "--enable-model" in captured[1] and "--disable-model" not in captured[1]


def test_live_scene_generation_requires_manual_gate(monkeypatch, tmp_path):
    from scripts.run_ppt_production import generate_scene_graph
    called = []
    monkeypatch.delenv("PPT_MODEL_LIVE_APPROVED", raising=False)
    monkeypatch.setattr("scripts.run_ppt_production.subprocess.run", lambda *_a, **_k: called.append(True))
    with pytest.raises(RuntimeError, match="PPT_MODEL_LIVE_APPROVED"):
        generate_scene_graph("never-run", tmp_path / "approved.md", tmp_path, tmp_path / "out", tmp_path, 20)
    assert not called
