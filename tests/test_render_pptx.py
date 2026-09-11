import json
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from src.render_pptx import RenderPptxError, build_slide_plan, build_slides_markdown, render_pptx, write_pptx_layout_audit


def _fixture(tmp_path: Path, *, with_image=True, sources=True) -> tuple[Path, Path | None]:
    assets = tmp_path / "assets"; assets.mkdir()
    if with_image: Image.new("RGB", (200, 100), "blue").save(assets / "system.png")
    image = "\n![系统架构](assets/system.png)\n图片来源：[来源: 手册.pdf, 第 10 页]\n" if with_image else ""
    markdown = tmp_path / "proposal.md"
    markdown.write_text("""# 园区智慧能源方案

## 总体架构

### 建设范围

- 配置 10 MW 光伏与 20 MWh 储能[来源: 手册.pdf, 第 10 页]
- 接入 10 kV 配电数据[来源: 规范.pdf, 第 7 页]

系统支持预测、决策、调度闭环，容量需现场确认。[来源: 手册.pdf, 第 10 页]
""" + image + """
## 实施计划

| 阶段 | 工作 | 周期 |
| --- | --- | --- |
| 一期 | 勘察 | 2 周 |

## 来源与依据

- 手册.pdf：第 10 页
""", encoding="utf-8")
    source = tmp_path / "proposal.sources.json"
    if sources: source.write_text(json.dumps({"sources": [{"source_id": "S1", "filename": "手册.pdf", "pages": [10], "chunk_id": "private"}, {"source_id": "S2", "filename": "规范.pdf", "pages": [7]}]}, ensure_ascii=False), encoding="utf-8")
    return markdown, source if sources else None


def test_plan_contains_cover_sections_tables_images_and_sources(tmp_path):
    markdown, _ = _fixture(tmp_path)
    plan, warnings = build_slide_plan(markdown, mode="presentation", max_slides=3)
    assert plan["slides"][0]["layout"] == "title"
    assert any(s["layout"] == "table" for s in plan["slides"])
    assert any(s["figure_ids"] for s in plan["slides"])
    assert plan["slides"][-1]["layout"] == "sources"
    assert plan["sources"]["S1"]["filename"] == "手册.pdf"
    assert not warnings or all("maximum" in x for x in warnings)


def test_render_valid_pptx_report_and_reopen(tmp_path):
    markdown, _ = _fixture(tmp_path)
    output = tmp_path / "proposal.pptx"
    report_path = render_pptx(markdown, output, mode="faithful")
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["errors"] == []
    with zipfile.ZipFile(output) as archive:
        assert "[Content_Types].xml" in archive.namelist()
        assert "ppt/presentation.xml" in archive.namelist()
        assert any(n.startswith("ppt/media/") for n in archive.namelist())
    assert len(Presentation(output).slides) == report["slide_count"]


def test_overflow_does_not_delete_content(tmp_path):
    markdown, _ = _fixture(tmp_path, with_image=False)
    markdown.write_text(markdown.read_text(encoding="utf-8").replace("系统支持预测、决策、调度闭环，容量需现场确认。", "。".join(f"第{i}项工程参数为 {i} kW" for i in range(18)) + "。"), encoding="utf-8")
    plan, warnings = build_slide_plan(markdown, max_slides=2)
    assert any("maximum" in warning for warning in warnings)
    assert any("（续）" in slide["title"] for slide in plan["slides"])
    assert "第17项工程参数为 17 kW" in "\n".join(sum((s["bullets"] for s in plan["slides"]), []))


def test_missing_sources_and_images_warn_without_http(tmp_path):
    markdown, _ = _fixture(tmp_path, with_image=False, sources=False)
    plan, warnings = build_slide_plan(markdown)
    assert not plan["sources"]
    assert plan["source_handoff"] == "LEGACY_INPUT_WARNING"
    assert any("LEGACY_INPUT_WARNING" in warning for warning in warnings)
    markdown.write_text(markdown.read_text(encoding="utf-8") + "\n![bad](assets/nope.png)\n", encoding="utf-8")
    _, warnings = build_slide_plan(markdown)
    assert any("skipped missing" in warning for warning in warnings)


def test_invalid_explicit_sources_json_warns(tmp_path):
    markdown, _ = _fixture(tmp_path, sources=False)
    invalid = tmp_path / "bad.json"; invalid.write_text("{", encoding="utf-8")
    _, warnings = build_slide_plan(markdown, sources_path=invalid)
    assert any("invalid sources JSON" in warning for warning in warnings)


def test_figures_keep_markdown_h2_when_no_sidecar_exists(tmp_path):
    assets = tmp_path / "assets"; assets.mkdir()
    for name in ("vpp.png", "arch.png", "price.png"):
        Image.new("RGB", (200, 100), "blue").save(assets / name)
    markdown = tmp_path / "proposal.md"
    markdown.write_text("""# 演示

## 系统总体架构

架构说明。[来源: 架构.pdf, 第 2 页]
![平台产品](assets/arch.png)
图片来源：架构.pdf，第 2 页
![虚拟电厂交易运营](assets/vpp.png)
图片来源：运营.pdf，第 3 页
![分时电价收益曲线](assets/price.png)
图片来源：收益.pdf，第 4 页

## 虚拟电厂调控与运营

调度说明。[来源: 运营.pdf, 第 3 页]

## 项目实施与综合效益

收益说明。[来源: 收益.pdf, 第 4 页]
""", encoding="utf-8")
    plan, _ = build_slide_plan(markdown)
    figures = plan["figures"]
    assert figures["F1"]["target_section"] == "系统总体架构"
    assert figures["F2"]["target_section"] == "系统总体架构"
    assert figures["F3"]["target_section"] == "系统总体架构"
    assert all(figure["visible_sources"] for figure in figures.values())
    figure_slides = [slide for slide in plan["slides"] if slide["figure_ids"]]
    assert len(figure_slides) == 3
    assert all(figure["original_section"] == "系统总体架构" for figure in figures.values())


def test_build_slides_writes_presentation_markdown_and_matching_plan(tmp_path):
    markdown, source = _fixture(tmp_path)
    output = tmp_path / "proposal.slides.md"
    slides, plan_path, warnings = build_slides_markdown(markdown, output, mode="presentation", max_slides=15, sources_path=source)
    text = slides.read_text(encoding="utf-8")
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    assert "## slide-001 | title |" in text
    assert "Figure F1:" in text
    assert "assets/system.png" not in text
    assert len(plan["slides"]) >= 4
    assert "chunk_id" not in plan["sources"]["S1"]
    assert isinstance(warnings, list)


def test_renderer_consumes_existing_slide_plan(tmp_path):
    markdown, source = _fixture(tmp_path)
    slides = tmp_path / "proposal.slides.md"
    _, plan_path, _ = build_slides_markdown(markdown, slides, sources_path=source)
    output = tmp_path / "proposal.pptx"
    report_path = render_pptx(markdown, output, plan_path=plan_path)
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["slide_plan"] == "proposal.slide-plan.json"
    assert report["source_handoff"] == "PASS"
    assert len(Presentation(output).slides) == len(json.loads(plan_path.read_text(encoding="utf-8"))["slides"])


def test_briefing_enforces_hard_cap_tracks_coverage_and_is_stable(tmp_path):
    markdown, _ = _fixture(tmp_path, with_image=False, sources=False)
    long_items = "\n".join(f"- 支撑说明 {index}，不影响核心指标。" for index in range(20))
    markdown.write_text(markdown.read_text(encoding="utf-8").replace("## 来源与依据", long_items + "\n\n## 来源与依据"), encoding="utf-8")
    first, _ = build_slide_plan(markdown, mode="briefing", max_slides=4)
    second, _ = build_slide_plan(markdown, mode="briefing", max_slides=4)
    assert len(first["slides"]) <= 4
    assert first["actual_slide_count"] == len(first["slides"])
    assert first["cap_enforced"] is True
    assert first["source_coverage"]
    assert first["slides"] == second["slides"]
    assert len({slide["slide_id"] for slide in first["slides"]}) == len(first["slides"])


def test_briefing_preserves_selected_numeric_threshold_and_uses_no_legacy_source_ids(tmp_path):
    markdown, _ = _fixture(tmp_path, with_image=False, sources=False)
    markdown.write_text(markdown.read_text(encoding="utf-8").replace("系统支持预测、决策、调度闭环，容量需现场确认。", "调节容量不低于 20 兆瓦，持续参与调峰不小于 2 小时，偏差不超过 15%。"), encoding="utf-8")
    plan, _ = build_slide_plan(markdown, mode="briefing", max_slides=8)
    content = "\n".join(bullet for slide in plan["slides"] for bullet in slide["bullets"])
    assert "20 兆瓦" in content and "2 小时" in content and "15%" in content
    assert plan["source_handoff"] == "LEGACY_INPUT_WARNING"
    assert not plan["sources"]
    assert all(not slide["source_ids"] for slide in plan["slides"])


def test_layout_audit_marks_visual_render_blocked_without_page_images(tmp_path):
    markdown, source = _fixture(tmp_path)
    slides = tmp_path / "proposal.slides.md"
    _, plan_path, _ = build_slides_markdown(markdown, slides, sources_path=source)
    pptx = tmp_path / "proposal.pptx"
    render_pptx(markdown, pptx, plan_path=plan_path)
    json_path, md_path = write_pptx_layout_audit(pptx, plan_path, tmp_path / "audit")
    audit = json.loads(json_path.read_text(encoding="utf-8"))
    assert md_path.is_file()
    assert audit["visual_render"] == "BLOCKED"
    assert audit["slide_count"] == audit["plan_slide_count"]
    assert all(page["shape_bounds_ok"] for page in audit["pages"])


def test_many_sources_fit_without_geometry_overflow(tmp_path):
    markdown, source = _fixture(tmp_path)
    records = [{"source_id": f"S{i}", "filename": f"来源文件{i}.pdf", "pages": [i]}
               for i in range(1, 40)]
    source.write_text(json.dumps({"sources": records}, ensure_ascii=False), encoding="utf-8")
    output = tmp_path / "many-sources.pptx"
    report = json.loads(render_pptx(markdown, output, sources_path=source).read_text(encoding="utf-8"))
    assert not report["errors"]


def test_briefing_uses_all_selected_figures_when_cap_has_space(tmp_path):
    assets = tmp_path / "assets"; assets.mkdir()
    for index in range(5): Image.new("RGB", (20, 10), "blue").save(assets / f"f{index}.png")
    markdown = tmp_path / "proposal.md"
    markdown.write_text("# T\n\n## A\n\n" + "\n\n".join(
        f"要点{index}，需结合现场确认。[来源: 手册.pdf, 第 1 页]\n![图{index}](assets/f{index}.png)"
        for index in range(5)) + "\n\n## 来源与依据\n- 手册.pdf：第 1 页\n", encoding="utf-8")
    source = tmp_path / "proposal.sources.json"; source.write_text(json.dumps({"sources": [{"source_id": "S1", "filename": "手册.pdf", "pages": [1]}]}, ensure_ascii=False), encoding="utf-8")
    plan, _ = build_slide_plan(markdown, mode="briefing", max_slides=10, sources_path=source)
    assert len(plan["slides"]) <= 10
    assert {fid for slide in plan["slides"] for fid in slide["figure_ids"]} == set(plan["figures"])
