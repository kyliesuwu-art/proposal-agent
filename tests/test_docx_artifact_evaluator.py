from docx import Document
from docx.shared import Mm
from src.artifact_evaluation.docx_evaluator import evaluate_docx
from src.artifact_evaluation.models import EvaluationStatus
from src.artifact_evaluation.profiles import get_profile
def test_docx_parse_a4_and_visual_not_run(tmp_path):
 p=tmp_path/"a.docx";d=Document();d.sections[0].page_width=Mm(210);d.sections[0].page_height=Mm(297);d.add_heading("Title",0);d.add_paragraph("body");d.save(p)
 r=evaluate_docx(p,get_profile("client-delivery"));assert r.status is EvaluationStatus.PASS;assert r.metrics["page_size"]["width_mm"]==210;assert r.metrics["visual_checks_status"]=="NOT_RUN"
def test_broken_docx_fails(tmp_path):
 p=tmp_path/"bad.docx";p.write_bytes(b"no");assert evaluate_docx(p,get_profile("client-delivery")).status is EvaluationStatus.FAIL
