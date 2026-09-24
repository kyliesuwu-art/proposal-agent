from __future__ import annotations

import asyncio
import hashlib
import json
from dataclasses import replace
from pathlib import Path

import pytest
from pptx import Presentation

from src.wecom.adapter import IncomingMessage, WeComAgentAdapter
from src.wecom.artifact_service import ArtifactService, PptResumeRejected
from src.wecom.models import ArtifactJob, ArtifactJobStatus, ArtifactType
from src.wecom.ppt_checkpoint import CheckpointStore
from src.wecom.ppt_failures import PptFailureCode, PptFailureInfo, PptRunnerFailure
from src.wecom.store import SQLiteTaskStore
from src.wecom.task_service import TaskService
from src.wecom.word_runner import WordResult


class ProposalRunner:
    def run(self, _request: str, output_dir: Path) -> Path:
        output_dir.mkdir(parents=True, exist_ok=True)
        proposal = output_dir / "proposal.md"
        proposal.write_text("# approved\n\n## confirmations\nowner【待确认】\n", encoding="utf-8")
        return proposal


class WordRunner:
    def run(self, _approved: Path, output_dir: Path, _task_id: str) -> WordResult:
        output_dir.mkdir(parents=True, exist_ok=True)
        path = output_dir / "proposal.docx"
        path.write_bytes(b"PK\x03\x04word")
        return WordResult(path)


class Transport:
    def __init__(self) -> None:
        self.files: list[tuple[str, Path]] = []
        self.messages: list[tuple[str, str]] = []

    async def reply_text(self, _frame: object, _text: str) -> None:
        return None

    async def send_text(self, chatid: str, text: str) -> None:
        self.messages.append((chatid, text))

    async def send_file(self, chatid: str, path: Path) -> None:
        self.files.append((chatid, path))


class FakePptRunner:
    def __init__(self, *, evaluation: str = "PASS_WITH_WARNINGS", failure: PptFailureInfo | None = None) -> None:
        self.evaluation = evaluation
        self.failure = failure
        self.calls: list[dict[str, object]] = []

    def run(
        self,
        approved_md_path: Path,
        task_root: Path,
        output_dir: Path,
        task_id: str,
        job_id: str,
        *,
        resume_from_job_root: Path | None = None,
    ):
        self.calls.append({
            "approved": approved_md_path,
            "task_root": task_root,
            "output_dir": output_dir,
            "task_id": task_id,
            "job_id": job_id,
            "resume_from_job_root": resume_from_job_root,
        })
        if self.failure:
            raise PptRunnerFailure(self.failure)
        output_dir.mkdir(parents=True, exist_ok=True)
        deck = output_dir / "proposal.pptx"
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        presentation.save(deck)
        from src.wecom.ppt_runner import PptResult
        return PptResult(deck, evaluation_status=self.evaluation, warnings=("visual_checks_not_run",) if self.evaluation == "PASS_WITH_WARNINGS" else ())


def _approved(tmp_path: Path, *, user: str = "u1", chat: str = "c1"):
    store = SQLiteTaskStore(tmp_path / "state.sqlite")
    tasks = TaskService(store, ProposalRunner(), output_root=tmp_path / "tasks")
    task, _ = tasks.receive(IncomingMessage(user, chat, "医院供配电改造方案，重点包括可靠性和智能巡检", "task-message"))
    tasks.wait_for_idle()
    task, _ = tasks.approve(task.task_id)
    return store, tasks, store.get(task.task_id)


def _source_job(store: SQLiteTaskStore, task, artifact_root: Path, *, pages: tuple[int, ...] = (1,), used: int = 3, artifact_type: ArtifactType = ArtifactType.PPTX) -> ArtifactJob:
    digest = hashlib.sha256(Path(task.approved_md_path).read_bytes()).hexdigest()
    root = artifact_root / task.task_id / "source-job"
    root.mkdir(parents=True)
    assets = root / "assets_manifest.json"
    assets.write_text('{"assets":[]}', encoding="utf-8")
    manifest_digest = hashlib.sha256(assets.read_bytes()).hexdigest()
    global_path = root / "global_art_direction.json"
    global_path.write_text('{"theme":"hospital"}', encoding="utf-8")
    checkpoints = CheckpointStore(
        root,
        task_id=task.task_id,
        job_id="source-job",
        approved_sha256=digest,
        assets_manifest_sha256=manifest_digest,
        compatibility={"producer": "v4", "prompt": "v4", "scene_schema": "v4", "pages": 20},
        model_max_calls=26,
    )
    checkpoints.record(stage="global_art_direction", canonical_path=global_path, page=None, calls_used=used)
    scene_root = root / "scene_graphs"
    for page in pages:
        scene = scene_root / f"page_{page:02}.json"
        scene.parent.mkdir(parents=True, exist_ok=True)
        scene.write_text(json.dumps({"page": page}), encoding="utf-8")
        checkpoints.record(stage="page_scene_graph", canonical_path=scene, page=page, calls_used=used)
    (root / "ark_calls.json").write_text(json.dumps([{"network_request_started": True} for _ in range(used)]), encoding="utf-8")
    now = "2026-09-24T00:00:00+08:00"
    source = ArtifactJob(
        "source-job", task.task_id, artifact_type, ArtifactJobStatus.FAILED,
        task.approved_md_path, digest, str(root), None, None, "page_design", "PPT generation failed", now, now,
        retryable=True, failure_code=PptFailureCode.MODEL_TRANSPORT_RETRY_EXHAUSTED.value,
    )
    store.create_artifact_job(source)
    return source


def _service(tmp_path: Path, runner: FakePptRunner):
    store, tasks, task = _approved(tmp_path)
    artifacts = ArtifactService(store, WordRunner(), runner, task_output_root=tmp_path / "tasks", output_root=tmp_path / "artifacts")
    transport = Transport()
    adapter = WeComAgentAdapter(tasks, transport, artifacts)
    return store, tasks, task, artifacts, transport, adapter


def _close(store: SQLiteTaskStore, tasks: TaskService, artifacts: ArtifactService) -> None:
    artifacts.shutdown()
    tasks.shutdown()
    store.close()


def test_resume_creates_child_runs_once_and_delivers_through_existing_event_flow(tmp_path: Path) -> None:
    runner = FakePptRunner()
    store, tasks, task, artifacts, transport, adapter = _service(tmp_path, runner)
    try:
        source = _source_job(store, task, tmp_path / "artifacts", pages=(1, 2, 4))
        before = {p.relative_to(Path(source.output_dir)): p.read_bytes() for p in Path(source.output_dir).rglob("*") if p.is_file()}

        async def scenario() -> ArtifactJob:
            adapter._loop = asyncio.get_running_loop()
            child = artifacts.resume_ppt(task_id=task.task_id, source_job_id=source.job_id, user_id="u1", chat_id="c1")
            await asyncio.to_thread(artifacts.wait_for_idle)
            await asyncio.sleep(0)
            return child

        child = asyncio.run(scenario())
        persisted = store.get_artifact_job(child.job_id)
        assert child.parent_job_id == source.job_id
        assert child.resume_from_job_id == source.job_id
        assert Path(child.output_dir) != Path(source.output_dir)
        assert len(runner.calls) == 1
        assert runner.calls[0]["resume_from_job_root"] == Path(source.output_dir)
        assert persisted.status is ArtifactJobStatus.DELIVERED
        assert len(transport.files) == 1
        assert store.get_artifact_job(source.job_id) == source
        assert before == {p.relative_to(Path(source.output_dir)): p.read_bytes() for p in Path(source.output_dir).rglob("*") if p.is_file()}
    finally:
        _close(store, tasks, artifacts)


def test_resume_evaluation_fail_never_sends_and_is_not_retryable(tmp_path: Path) -> None:
    runner = FakePptRunner(evaluation="FAIL")
    store, tasks, task, artifacts, transport, _adapter = _service(tmp_path, runner)
    try:
        source = _source_job(store, task, tmp_path / "artifacts")
        child = artifacts.resume_ppt(task_id=task.task_id, source_job_id=source.job_id, user_id="u1", chat_id="c1")
        artifacts.wait_for_idle()
        failed = store.get_artifact_job(child.job_id)
        assert failed.status is ArtifactJobStatus.FAILED
        assert failed.failure_code == PptFailureCode.EVALUATION_FAILED.value
        assert failed.retryable is False
        assert not transport.files
    finally:
        _close(store, tasks, artifacts)


@pytest.mark.parametrize(
    "failure",
    [
        PptFailureInfo(PptFailureCode.MODEL_TRANSPORT_RETRY_EXHAUSTED, "page_design:10", True, "PPT model transport attempts were exhausted."),
        PptFailureInfo(PptFailureCode.MODEL_PROVIDER_TRANSIENT, "page_design:10", True, "PPT model provider attempts were exhausted."),
        PptFailureInfo(PptFailureCode.MODEL_OUTPUT_CONTRACT_RETRY_EXHAUSTED, "page_design:10", True, "PPT model output contract attempts were exhausted."),
        PptFailureInfo(PptFailureCode.MODEL_BUDGET_EXHAUSTED, "page_design:10", False, "PPT model call budget was exhausted."),
        PptFailureInfo(PptFailureCode.CONFIG_INVALID, "startup", False, "PPT producer configuration is invalid."),
        PptFailureInfo(PptFailureCode.AUTH_REJECTED, "startup", False, "PPT provider credentials were rejected."),
        PptFailureInfo(PptFailureCode.SECURITY_REJECTED, "resume_validation", False, "PPT producer security validation failed."),
        PptFailureInfo(PptFailureCode.INPUT_INVALID, "startup", False, "PPT producer input is invalid."),
        PptFailureInfo(PptFailureCode.RENDER_FAILED, "render", False, "PPT rendering failed."),
    ],
)
def test_resume_persists_structured_runner_failure_without_message_parsing(tmp_path: Path, failure: PptFailureInfo) -> None:
    runner = FakePptRunner(failure=failure)
    store, tasks, task, artifacts, _transport, _adapter = _service(tmp_path, runner)
    try:
        source = _source_job(store, task, tmp_path / "artifacts")
        child = artifacts.resume_ppt(task_id=task.task_id, source_job_id=source.job_id, user_id="u1", chat_id="c1")
        artifacts.wait_for_idle()
        failed = store.get_artifact_job(child.job_id)
        assert (failed.failure_code, failed.error_stage, failed.retryable) == (failure.code.value, failure.stage, failure.retryable)
        assert failed.error_message == "PPT generation failed"
    finally:
        _close(store, tasks, artifacts)


def test_resume_rejects_invalid_source_without_creating_job_or_runner_call(tmp_path: Path) -> None:
    runner = FakePptRunner()
    store, tasks, task, artifacts, _transport, _adapter = _service(tmp_path, runner)
    try:
        source = _source_job(store, task, tmp_path / "artifacts")
        Path(source.output_dir, "assets_manifest.json").write_text('{"changed":true}', encoding="utf-8")
        with pytest.raises(PptResumeRejected):
            artifacts.resume_ppt(task_id=task.task_id, source_job_id=source.job_id, user_id="u1", chat_id="c1")
        assert store.artifact_jobs_for(task.task_id, ArtifactType.PPTX) == [source]
        assert not runner.calls
    finally:
        _close(store, tasks, artifacts)


@pytest.mark.parametrize("field, value", [
    ("status", ArtifactJobStatus.READY),
    ("retryable", False),
    ("artifact_type", ArtifactType.WORD),
])
def test_resume_rejects_ineligible_source_identity(tmp_path: Path, field: str, value: object) -> None:
    runner = FakePptRunner()
    store, tasks, task, artifacts, _transport, _adapter = _service(tmp_path, runner)
    try:
        source = _source_job(
            store, task, tmp_path / "artifacts",
            artifact_type=value if field == "artifact_type" else ArtifactType.PPTX,
        )
        if field != "artifact_type":
            store.update_artifact_job(replace(source, **{field: value}))
        with pytest.raises(PptResumeRejected):
            artifacts.resume_ppt(task_id=task.task_id, source_job_id=source.job_id, user_id="u1", chat_id="c1")
        assert not runner.calls
    finally:
        _close(store, tasks, artifacts)


def test_resume_rejects_missing_checkpoint_and_chat_mismatch(tmp_path: Path) -> None:
    runner = FakePptRunner()
    store, tasks, task, artifacts, _transport, _adapter = _service(tmp_path, runner)
    try:
        source = _source_job(store, task, tmp_path / "artifacts")
        with pytest.raises(PptResumeRejected):
            artifacts.resume_ppt(task_id=task.task_id, source_job_id=source.job_id, user_id="u1", chat_id="other-chat")
        Path(source.output_dir, "ppt_generation_checkpoints.json").unlink()
        with pytest.raises(PptResumeRejected):
            artifacts.resume_ppt(task_id=task.task_id, source_job_id=source.job_id, user_id="u1", chat_id="c1")
        assert not runner.calls
    finally:
        _close(store, tasks, artifacts)

def test_resume_rejects_user_or_chat_mismatch_and_exhausted_budget(tmp_path: Path) -> None:
    runner = FakePptRunner()
    store, tasks, task, artifacts, _transport, _adapter = _service(tmp_path, runner)
    try:
        source = _source_job(store, task, tmp_path / "artifacts", used=26)
        with pytest.raises(PptResumeRejected):
            artifacts.resume_ppt(task_id=task.task_id, source_job_id=source.job_id, user_id="other", chat_id="c1")
        with pytest.raises(PptResumeRejected):
            artifacts.resume_ppt(task_id=task.task_id, source_job_id=source.job_id, user_id="u1", chat_id="c1")
        assert not runner.calls
    finally:
        _close(store, tasks, artifacts)

def test_resume_rejects_active_child_and_preserves_no_second_output_dir(tmp_path: Path) -> None:
    runner = FakePptRunner()
    store, tasks, task, artifacts, _transport, _adapter = _service(tmp_path, runner)
    try:
        source = _source_job(store, task, tmp_path / "artifacts")
        now = "2026-09-24T00:00:00+08:00"
        active = ArtifactJob(
            "active-child", task.task_id, ArtifactType.PPTX, ArtifactJobStatus.QUEUED,
            source.source_md_path, source.source_md_sha256, str(tmp_path / "artifacts" / task.task_id / "active-child"),
            None, None, None, None, now, now, parent_job_id=source.job_id, resume_from_job_id=source.job_id,
        )
        store.create_resume_child_atomic(source.job_id, active)
        with pytest.raises(PptResumeRejected):
            artifacts.resume_ppt(task_id=task.task_id, source_job_id=source.job_id, user_id="u1", chat_id="c1")
        assert len(store.list_artifact_children(source.job_id)) == 1
        assert not Path(active.output_dir).exists()
        assert not runner.calls
    finally:
        _close(store, tasks, artifacts)


def test_resume_unknown_exception_is_conservative_and_child_lineage_persists_after_reopen(tmp_path: Path) -> None:
    class BrokenRunner(FakePptRunner):
        def run(self, *args, **kwargs):
            raise RuntimeError("Timeout wording must not decide retryability")

    runner = BrokenRunner()
    store, tasks, task, artifacts, _transport, _adapter = _service(tmp_path, runner)
    try:
        source = _source_job(store, task, tmp_path / "artifacts")
        child = artifacts.resume_ppt(task_id=task.task_id, source_job_id=source.job_id, user_id="u1", chat_id="c1")
        artifacts.wait_for_idle()
        failed = store.get_artifact_job(child.job_id)
        assert failed.failure_code == PptFailureCode.UNKNOWN_FAILURE.value
        assert failed.retryable is False
        artifacts.shutdown()
        reopened = SQLiteTaskStore(tmp_path / "state.sqlite")
        saved = reopened.get_artifact_job(child.job_id)
        assert (saved.parent_job_id, saved.resume_from_job_id) == (source.job_id, source.job_id)
        reopened.close()
    finally:
        tasks.shutdown()
        store.close()
