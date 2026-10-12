from pptx import Presentation
from pptx.util import Inches,Pt
from src.artifact_evaluation.pptx_evaluator import evaluate_pptx
from src.artifact_evaluation.models import EvaluationStatus
from src.artifact_evaluation.profiles import get_profile
def test_pptx_parse_and_bounds(tmp_path):
 p=tmp_path/"a.pptx";prs=Presentation();s=prs.slides.add_slide(prs.slide_layouts[6]);box=s.shapes.add_textbox(Inches(0),Inches(0),Inches(4),Inches(1));box.text_frame.paragraphs[0].add_run().text="1 Title";prs.save(p)
 r=evaluate_pptx(p,get_profile("client-delivery"));assert r.metrics["slide_count"]==1;assert r.metrics["visual_checks_status"]=="NOT_RUN"
 p2=tmp_path/"bad.pptx";p2.write_bytes(b"no");assert evaluate_pptx(p2,get_profile("client-delivery")).status is EvaluationStatus.FAIL


def test_client_delivery_rejects_unremoved_internal_and_source_leaks(tmp_path):
 p=tmp_path/"leak.pptx";prs=Presentation();s=prs.slides.add_slide(prs.slide_layouts[6]);box=s.shapes.add_textbox(Inches(0),Inches(0),Inches(4),Inches(1));box.text="1 【待确认】 来源：设计母稿.pdf，第10页";prs.save(p)
 r=evaluate_pptx(p,get_profile("client-delivery"));assert r.status is EvaluationStatus.FAIL;assert r.metrics["internal_label_count"]>=1;assert r.metrics["source_leak_count"]==1
