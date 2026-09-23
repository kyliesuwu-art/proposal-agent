"""Ark-led eight-page visual calibration, resumable and isolated from V2/V3.

Ark owns the page composition.  This runner only supplies locked facts and
rendered visual references, persists every response, validates the returned
Scene Graphs, applies the existing common chrome, and renders with PowerPoint.
"""
from __future__ import annotations

import base64
import argparse
import http.client
import importlib.util
import json
import os
import re
import socket
import shutil
import subprocess
import sys
import time
import urllib.error
from copy import deepcopy
from pathlib import Path
from typing import Callable

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent.parent
FULL15 = ROOT / "outputs" / "ppt_pure_art_director" / "full15"
BASE = FULL15 / "mainline_expanded_v2"
V3 = FULL15 / "visual_calibration_v3"
OUT = FULL15 / "visual_calibration_v4_ark"
PICK = (1, 2, 3, 5, 6, 14, 19, 20)
NAMES = {1: "G", 2: "TOC", 3: "BACKGROUND", 5: "H", 6: "ARCH", 14: "D", 19: "E", 20: "CLOSE"}
APPROVED_TEXT = ""
# Production injects these from the already-built task assets manifest.  The
# references attached to the model request are deliberately not included.
CONTENT_ASSETS: tuple[dict[str, object], ...] = ()
V3_PAGE_MAP = {1: 1, 2: 2, 3: 3, 5: 4, 6: 5, 14: 6, 19: 7, 20: 8}
REFERENCE_SHEETS = (
    ROOT / "outputs" / "ppt_template_library_v1" / "source_analysis" / "comking_company" / "contact_sheet.png",
    ROOT / "outputs" / "ppt_template_library_v1" / "source_analysis" / "fujian" / "contact_sheet.png",
    ROOT / "outputs" / "ppt_template_library_v1" / "source_analysis" / "industrial_park" / "contact_sheet.png",
    ROOT / "outputs" / "ppt_template_library_v1" / "source_analysis" / "jiamusi" / "contact_sheet.png",
)


class ModelCallBudget:
    """A producer-wide metered-call budget consumed before every Ark request."""
    def __init__(self, maximum: int) -> None:
        if maximum <= 0:
            raise ValueError("model call budget must be positive")
        self.maximum, self.used = maximum, 0

    def consume(self, _call_id: str) -> None:
        if self.used >= self.maximum:
            raise RuntimeError("PPT model call budget exhausted before Ark request")
        self.used += 1


MODEL_CALL_BUDGET: ModelCallBudget | None = None


class ModelJsonContractError(ValueError):
    """A model response could not satisfy the JSON output contract."""

    def __init__(self, message: str, *, line: int | None = None, column: int | None = None) -> None:
        super().__init__(message)
        self.line, self.column = line, column


class ModelTransportError(RuntimeError):
    """A model HTTP attempt failed before a complete response was available."""

    def __init__(self, message: str, *, retryable: bool) -> None:
        super().__init__(message)
        self.retryable = retryable


def configure_model_call_budget(maximum: int | None) -> None:
    global MODEL_CALL_BUDGET
    MODEL_CALL_BUDGET = ModelCallBudget(maximum) if maximum is not None else None


def load(name: str, filename: str):
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / filename)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader
    spec.loader.exec_module(module)
    return module


PAD = load("pad_v4", "run_ppt_pure_art_director.py")
EDITORIAL = load("editorial_v4", "run_ppt_pure_art_director_editorial.py")


def say(message: str) -> None:
    print(f"[{time.strftime('%H:%M')}] {message}", flush=True)


def page_png(folder: Path, page: int) -> Path:
    matches = [p for p in (folder / "pages" / "slide").glob("*.png") if re.search(r"(\d+)$", p.stem) and int(re.search(r"(\d+)$", p.stem).group(1)) == page]
    if len(matches) != 1:
        raise RuntimeError(f"could not locate rendered page {page} in {folder}")
    return matches[0]


def image_url(path: Path, limit: int = 1100) -> str:
    """Use a bounded reference image, keeping each Ark request practical."""
    with Image.open(path) as source:
        image = source.convert("RGB")
        if max(image.size) > limit:
            scale = limit / max(image.size)
            image = image.resize((round(image.width * scale), round(image.height * scale)), Image.Resampling.LANCZOS)
        from io import BytesIO
        stream = BytesIO(); image.save(stream, format="JPEG", quality=86, optimize=True)
    return "data:image/jpeg;base64," + base64.b64encode(stream.getvalue()).decode("ascii")


def image_part(path: Path, label: str) -> list[dict]:
    return [{"type": "text", "text": label}, {"type": "image_url", "image_url": {"url": image_url(path)}}]


def source_scene(page: int) -> dict:
    if APPROVED_TEXT:
        return {"elements": []}
    path = BASE / "scene_graphs" / f"page_{page:02}_{NAMES[page]}.json"
    return json.loads(path.read_text(encoding="utf-8"))


def text_content(scene: dict) -> list[str]:
    return [str(e["text"]) for e in scene.get("elements", []) if e.get("type") == "text" and str(e.get("text", "")).strip()]


def allowed_assets(scene: dict) -> list[str]:
    return sorted({str(e["image_source"]) for e in scene.get("elements", []) if e.get("type") == "image" and str(e.get("image_source", "")).startswith("assets/")})


def context_for(page: int) -> dict:
    scene = source_scene(page)
    allowed_content_images = (
        [dict(asset) for asset in CONTENT_ASSETS]
        if APPROVED_TEXT
        else [{"image_source": source, "usage": "content_asset"} for source in allowed_assets(scene)]
    )
    return {
        "page": page,
        "page_name": NAMES[page],
        "locked_display_text": text_content(scene),
        "allowed_content_images": allowed_content_images,
        "allowed_image_sources": [
            str(reference)
            for asset in allowed_content_images
            for reference in (asset.get("asset_id"), asset["image_source"])
            if reference
        ],
        "known_constraint": "Image image-001.jpg is only 537x489; do not make it a full-bleed or enlarged hero image." if page == 5 else None,
        "approved_markdown": APPROVED_TEXT if APPROVED_TEXT else None,
    }


def record_call(item: dict) -> None:
    path = OUT / "ark_calls.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    calls = json.loads(path.read_text(encoding="utf-8")) if path.is_file() else []
    calls.append(item); path.write_text(json.dumps(calls, ensure_ascii=False, indent=2), encoding="utf-8")


def _redact_error_message(error: BaseException) -> str:
    """Keep telemetry useful without retaining credentials or signed URLs."""
    message = str(error)
    for name in ("ARK_API_KEY", "DASHSCOPE_API_KEY", "AIBOT_SECRET"):
        value = os.environ.get(name)
        if value:
            message = message.replace(value, "[REDACTED]")
    message = re.sub(r"(?i)(authorization\s*[:=]\s*bearer\s+)[^\s]+", r"\1[REDACTED]", message)
    return re.sub(r"(https?://[^\s?]+)\?[^\s]+", r"\1?[REDACTED]", message)


def is_retryable_transport_error(error: BaseException) -> bool:
    """Classify only transient transport/status failures used by this client."""
    if isinstance(error, urllib.error.HTTPError):
        return error.code in {408, 429} or 500 <= error.code <= 599
    if isinstance(error, (http.client.IncompleteRead, ConnectionResetError, ConnectionAbortedError,
                          BrokenPipeError, TimeoutError, socket.timeout)):
        return True
    if isinstance(error, urllib.error.URLError):
        return isinstance(error.reason, (ConnectionResetError, ConnectionAbortedError, BrokenPipeError,
                                         TimeoutError, socket.timeout, OSError))
    try:
        import httpx  # type: ignore
        if isinstance(error, (httpx.ReadError, httpx.ReadTimeout, httpx.RemoteProtocolError)):
            return True
    except ImportError:
        pass
    try:
        import openai  # type: ignore
        if isinstance(error, (openai.APIConnectionError, openai.APITimeoutError)):
            return True
        if isinstance(error, openai.APIStatusError):
            return error.status_code in {408, 429} or 500 <= error.status_code <= 599
    except ImportError:
        pass
    return False


def _unique_top_level_object(payload: str) -> str:
    """Return one balanced object without rewriting any model-provided text."""
    depth = 0
    in_string = False
    escaped = False
    found: tuple[int, int] | None = None
    for index, char in enumerate(payload):
        if in_string:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                in_string = False
            continue
        if char == '"':
            in_string = True
        elif char == "{":
            if depth == 0:
                if found is not None:
                    raise ModelJsonContractError("model response contains multiple top-level JSON objects")
                start = index
            depth += 1
        elif char == "}" and depth:
            depth -= 1
            if depth == 0:
                found = (start, index + 1)
    if found is None or depth or in_string:
        return payload
    return payload[found[0]:found[1]]


def parse_model_json(raw: str) -> dict:
    """Apply only lossless JSON transport normalization and parse one object."""
    payload = raw.removeprefix("\ufeff").strip()
    lines = payload.splitlines()
    fence_lines = [index for index, line in enumerate(lines) if line.strip().startswith("```")]
    if fence_lines:
        if len(fence_lines) != 2 or fence_lines[0] != 0 or fence_lines[1] != len(lines) - 1:
            raise ModelJsonContractError("model response contains an ambiguous code fence")
        opener, closer = lines[0].strip().lower(), lines[-1].strip()
        if opener not in {"```", "```json"} or closer != "```":
            raise ModelJsonContractError("model response contains an unsupported code fence")
        payload = "\n".join(lines[1:-1]).strip()
    payload = _unique_top_level_object(payload)
    try:
        value = json.loads(payload)
    except json.JSONDecodeError as exc:
        raise ModelJsonContractError(f"invalid model JSON at line {exc.lineno}, column {exc.colno}: {exc.msg}", line=exc.lineno, column=exc.colno) from exc
    if not isinstance(value, dict):
        raise ModelJsonContractError("model JSON root is not an object")
    return value


GLOBAL_ART_DIRECTION_FIELDS = (
    "design_intent", "visual_personality", "color_system", "typography_hierarchy",
    "spacing_rhythm", "image_treatment", "diagram_language", "brand_motif",
    "page_families", "density_strategy", "do_not_use", "page_directions",
)


def validate_global_art_direction(value: dict) -> None:
    """Validate the explicitly prompted Global Art Direction contract."""
    missing = [field for field in GLOBAL_ART_DIRECTION_FIELDS if field not in value]
    if missing:
        raise ModelJsonContractError("global art direction missing required fields: " + ", ".join(missing))
    text_fields = {"design_intent", "visual_personality", "typography_hierarchy", "spacing_rhythm", "image_treatment", "diagram_language", "brand_motif", "density_strategy"}
    invalid_text = sorted(field for field in text_fields if not isinstance(value[field], str) or not value[field].strip())
    if invalid_text:
        raise ModelJsonContractError("global art direction requires non-empty text fields: " + ", ".join(invalid_text))
    invalid_objects = sorted(field for field in {"color_system", "page_families", "page_directions"} if not isinstance(value[field], dict) or not value[field])
    if invalid_objects:
        raise ModelJsonContractError("global art direction requires non-empty object fields: " + ", ".join(invalid_objects))
    if not isinstance(value["do_not_use"], list) or not value["do_not_use"] or not all(isinstance(item, str) and item.strip() for item in value["do_not_use"]):
        raise ModelJsonContractError("global art direction requires a non-empty do_not_use string list")

def _attempt_raw_file(raw_file: Path, attempt: int) -> Path:
    stem = raw_file.stem.removesuffix("_raw")
    return raw_file.with_name(f"{stem}_attempt_{attempt:02}_raw{raw_file.suffix}")


def _retry_system_prompt(system: str, error: BaseException) -> str:
    if isinstance(error, ModelJsonContractError):
        location = f" at line {error.line}, column {error.column}" if error.line else ""
        return system + "\nYour previous response was not valid JSON or Scene Graph validation" + location + ": " + str(error) + ". Return only one complete JSON object. For an asset whitelist failure, use the exact invalid_image_sources and allowed_image_sources above; style references are not content assets. If none fit, remove the image node."
    return system + "\nThe previous request ended before a complete response was received. Return only one complete JSON object with all required fields and no markdown, commentary, or code fence."


def _attempt_record(*, call_id: str, attempt: int, started: float, meta: dict, content: list[dict],
                    network_request_started: bool, http_success: bool, response_complete: bool,
                    parse_success: bool, schema_success: bool, success: bool, retryable: bool,
                    error: BaseException | None, response_path: Path | None) -> dict:
    return {
        "stage": call_id,
        "logical_call_id": call_id,
        "call_id": call_id,
        "attempt": attempt,
        **meta,
        "started_at": started,
        "finished_at": time.time(),
        "duration_seconds": round(time.time() - started, 3),
        "input_image_count": sum(1 for item in content if item.get("type") == "image_url"),
        "network_request_started": network_request_started,
        "http_success": http_success,
        "response_complete": response_complete,
        "raw_response_complete": response_complete,
        "parse_success": parse_success,
        "schema_success": schema_success,
        "success": success,
        "retryable": retryable,
        "response_path": str(response_path.resolve().relative_to(OUT.resolve())) if response_path else None,
        "error_type": type(error).__name__ if error else None,
        "error_message": _redact_error_message(error) if error else None,
        "error": _redact_error_message(error) if error else None,
        "json_line": error.line if isinstance(error, ModelJsonContractError) else None,
        "json_column": error.column if isinstance(error, ModelJsonContractError) else None,
        "budget_used": MODEL_CALL_BUDGET.used if MODEL_CALL_BUDGET else None,
        "budget_limit": MODEL_CALL_BUDGET.maximum if MODEL_CALL_BUDGET else None,
    }


def ask_json(*, call_id: str, system: str, content: list[dict], raw_file: Path, max_tokens: int = 12000,
             validator: Callable[[dict], None] | None = None, retry_delay_seconds: float = 2.0,
             sleep_fn: Callable[[float], None] = time.sleep) -> tuple[dict, dict]:
    """Perform at most two total budgeted HTTP attempts for one JSON stage."""
    last_error: BaseException | None = None
    for attempt in (1, 2):
        started = time.time()
        meta: dict = {"purpose": call_id}
        network_request_started = http_success = response_complete = parse_success = schema_success = False
        response_path: Path | None = None
        error: BaseException | None = None
        retryable = False
        try:
            if MODEL_CALL_BUDGET is not None:
                MODEL_CALL_BUDGET.consume(call_id)
            network_request_started = True
            say(f"Ark {call_id} attempt {attempt} request started")
            active_system = system if last_error is None else _retry_system_prompt(system, last_error)
            raw, meta = PAD.ask([{"role": "system", "content": active_system}, {"role": "user", "content": content}], call_id, max_tokens=max_tokens)
            http_success = response_complete = True
            response_path = _attempt_raw_file(raw_file, attempt)
            response_path.parent.mkdir(parents=True, exist_ok=True)
            response_path.write_text(raw, encoding="utf-8")
            parsed = parse_model_json(raw)
            parse_success = True
            if validator:
                validator(parsed)
            schema_success = True
        except ModelJsonContractError as exc:
            error = exc
            retryable = True
        except Exception as exc:
            error = exc
            telemetry = getattr(exc, "telemetry", None)
            if isinstance(telemetry, dict):
                meta = telemetry
            retryable = is_retryable_transport_error(exc)
        else:
            record_call(_attempt_record(call_id=call_id, attempt=attempt, started=started, meta=meta, content=content,
                                        network_request_started=network_request_started, http_success=http_success,
                                        response_complete=response_complete, parse_success=parse_success,
                                        schema_success=schema_success, success=True, retryable=False, error=None,
                                        response_path=response_path))
            say(f"Ark {call_id} attempt {attempt} saved ({meta.get('elapsed_seconds')}s; ttft={meta.get('ttft_seconds')}s)")
            return parsed, meta

        record_call(_attempt_record(call_id=call_id, attempt=attempt, started=started, meta=meta, content=content,
                                    network_request_started=network_request_started, http_success=http_success,
                                    response_complete=response_complete, parse_success=parse_success,
                                    schema_success=schema_success, success=False, retryable=retryable, error=error,
                                    response_path=response_path))
        if not retryable or attempt == 2:
            if isinstance(error, ModelJsonContractError):
                raise ModelJsonContractError(f"{call_id} failed after {attempt} HTTP attempts: {error}", line=error.line, column=error.column) from error
            raise ModelTransportError(f"{call_id} transport failure after {attempt} HTTP attempts: {type(error).__name__}: {_redact_error_message(error)}", retryable=retryable) from error
        last_error = error
        if retry_delay_seconds:
            sleep_fn(retry_delay_seconds)
    raise AssertionError("unreachable")

def art_direction() -> dict:
    target = OUT / "global_art_direction.json"
    if target.is_file():
        say("global art direction checkpoint reused")
        return json.loads(target.read_text(encoding="utf-8"))
    content = [{"type": "text", "text": "You are the global art director for an enterprise COMKING hospital smart-energy proposal. The following are real visual references: four company decks, then V2 and V3 calibration contact sheets. Analyze them visually, not by assuming textual descriptions."}]
    labels = ("COMKING company deck contact sheet", "Fujian nuclear smart-building deck contact sheet", "industrial-park smart-energy deck contact sheet", "Jiamusi university comprehensive-energy deck contact sheet")
    for label, sheet in zip(labels, REFERENCE_SHEETS): content.extend(image_part(sheet, label))
    content.extend(image_part(BASE / "contact_sheet.png", "Current complete V2 contact sheet"))
    content.extend(image_part(V3 / "contact_sheet.png", "Local V3 eight-page calibration contact sheet"))
    content.append({"type": "text", "text": json.dumps({"target_pages": [context_for(x) for x in PICK], "strict": ["Use mature corporate proposal design, not web UI", "Avoid repeated rounded-card grids, pill labels, meaningless diagonal or crossing lines", "Make diagrams communicate relationships", "Keep all text editable and preserve facts/boundaries", "Use one official logo and renderer-owned footer only"], "required_keys": ["design_intent", "visual_personality", "color_system", "typography_hierarchy", "spacing_rhythm", "image_treatment", "diagram_language", "brand_motif", "page_families", "density_strategy", "do_not_use", "page_directions"], "field_contract": {"design_intent": "non-empty string", "visual_personality": "non-empty string", "color_system": "non-empty object", "typography_hierarchy": "non-empty string", "spacing_rhythm": "non-empty string", "image_treatment": "non-empty string", "diagram_language": "non-empty string", "brand_motif": "non-empty string", "page_families": "non-empty object", "density_strategy": "non-empty string", "do_not_use": "non-empty array of strings", "page_directions": "non-empty object"}}, ensure_ascii=False)})
    system = "Return only a valid JSON object. You are judging actual attached images. Do not output a slide Scene Graph in this call; output the requested Global Art Direction object with concise, actionable page directions for pages 1,2,3,5,6,14,19,20."
    value, _ = ask_json(call_id="global_art_direction", system=system, content=content, raw_file=OUT / "global_art_direction_raw.json", max_tokens=7000, validator=validate_global_art_direction)
    target.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    return value


def page_request(page: int, direction: dict, *, revision: bool = False) -> dict:
    scene_path = OUT / "scene_graphs" / f"page_{page:02}_{NAMES[page]}.json"
    if scene_path.is_file() and not revision:
        say(f"page {page:02} scene checkpoint reused")
        return json.loads(scene_path.read_text(encoding="utf-8"))
    ctx = context_for(page)
    content = [{"type": "text", "text": "Attached images are STYLE_REFERENCE_ONLY, NOT_A_CONTENT_ASSET, and MUST_NOT_BE_REFERENCED_IN_SCENE_GRAPH. Use them only as visual baselines, but create a materially more mature composition based on Global Art Direction."}]
    content.extend(image_part(page_png(BASE, page), f"V2 page {page:02}"))
    if page in V3_PAGE_MAP:
        content.extend(image_part(page_png(V3, V3_PAGE_MAP[page]), f"V3 counterpart for original page {page:02}"))
    else:
        content.extend(image_part(V3 / "contact_sheet.png", "V3 calibration contact sheet: use as family reference, not as this page's layout"))
    content.extend(image_part(REFERENCE_SHEETS[0], "Closest real COMKING reference deck contact sheet"))
    content.append({"type": "text", "text": json.dumps({"global_art_direction": direction, "page_context": ctx, "canvas_inches": [13.333, 7.5], "safe_zones": {"logo": [0.45, 0.08, 2.05, 0.48], "footer": [0.55, 6.84, 12.15, 0.48]}, "schema": {"root": ["background", "elements"], "element_types": ["text", "image", "rectangle", "rounded_rectangle", "circle", "line", "arrow"], "common": ["id", "type", "x", "y", "w", "h", "z_order"], "text": ["text", "font_size", "font_weight", "alignment", "color"], "image": ["image_source", "crop"], "line_or_arrow": ["from", "to", "color", "line_width"]}, "requirements": ["Return a complete parseable Scene Graph only", "Do not draw logo/footer/page number", "Keep all main content outside safe zones", "Use only allowed_content_images as image_source. Node id is an element name, not an asset. STYLE_REFERENCE_ONLY attachments are NOT_A_CONTENT_ASSET and MUST_NOT_BE_REFERENCED_IN_SCENE_GRAPH. If none fit, omit the image node.", "All coordinates within canvas", "Text in normal body >=16pt; no markdown, URLs or new facts", "Do not rely on a local template; choose a page-specific composition", "The result must be directly editable with native PowerPoint shapes"]}, ensure_ascii=False)})
    if revision:
        previous = json.loads(scene_path.read_text(encoding="utf-8")); content.append({"type": "text", "text": "This is a single targeted revision after visual critique. Preserve locked facts and improve only the cited defects. Current Scene Graph: " + json.dumps(previous, ensure_ascii=False)})
    system = "You are a senior enterprise presentation art director. You can see the attached actual images. Return ONLY one JSON object Scene Graph, with background and elements. Do not explain your work. Avoid card-grid/UI conventions; every line, arrow, color block, or image crop must carry information or hierarchy."
    folder = OUT / ("revisions" if revision else "raw")
    raw = folder / f"page_{page:02}_{NAMES[page]}_{'revision_' if revision else ''}raw.json"
    def validate_scene_graph(value: dict) -> None:
        allowed_image_sources = set(ctx["allowed_image_sources"])
        validation = PAD.validate(value, allowed_image_sources)
        if validation["hard_errors"] or validation["preclamp_out_of_bounds"]:
            invalid_image_sources = sorted({
                str(element.get("image_source"))
                for element in value.get("elements", [])
                if isinstance(element, dict)
                and element.get("type") == "image"
                and element.get("image_source") not in allowed_image_sources
            })
            detail = (
                f" invalid_image_sources={invalid_image_sources}; "
                f"allowed_image_sources={sorted(allowed_image_sources)}; "
                "STYLE_REFERENCE_ONLY attachments cannot be used as image_source."
                if invalid_image_sources else ""
            )
            raise ModelJsonContractError(f"page {page:02} invalid Ark scene:{detail} validation={validation}")

    value, _ = ask_json(call_id=f"page_{page:02}_{'revision' if revision else 'design'}", system=system, content=content, raw_file=raw, validator=validate_scene_graph)
    scene_path.parent.mkdir(parents=True, exist_ok=True); scene_path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    return value


def draw(scenes: dict[int, dict]) -> Path:
    logo = EDITORIAL.extract_formal_logo(OUT / "assets")
    prs = Presentation(); prs.slide_width = Inches(PAD.W); prs.slide_height = Inches(PAD.H); blank = prs.slide_layouts[6]
    for page in PICK:
        slide = prs.slides.add_slide(blank)
        EDITORIAL.chrome(slide, page, logo_path=logo, total_pages=20)
        positions = {}
        for raw in sorted(scenes[page]["elements"], key=lambda e: float(e.get("z_order", e.get("z", 10)))):
            element = PAD.clamp(dict(raw)); typ = element.get("type")
            if typ in {"text", "number", "caption"}: shape = PAD.add_text(slide, element)
            elif typ == "image": shape = PAD.add_image(slide, element)
            elif typ in {"rectangle", "rounded_rectangle", "circle", "callout"}: shape = PAD.add_shape(slide, element)
            elif typ in {"line", "arrow"}: shape = PAD.add_line(slide, element, positions)
            else: continue
            positions[str(element["id"])] = shape
    deck = OUT / "pure_art_director_visual_calibration_v4_ark.pptx"; prs.save(deck); return deck


def render(deck: Path) -> list[Path]:
    pages = OUT / "pages"; pages.mkdir(parents=True, exist_ok=True)
    subprocess.run(["powershell", "-ExecutionPolicy", "Bypass", "-File", str(ROOT / "scripts" / "render_pptx_windows.ps1"), "-InputPath", str(deck), "-OutputDir", str(pages)], cwd=ROOT, check=True, timeout=420)
    cmd = "$a=New-Object -ComObject PowerPoint.Application;$a.Visible=$true;$p=$a.Presentations.Open('" + str(deck.resolve()).replace("'", "''") + "',$true,$true,$false);$p.SaveAs('" + str((OUT / "pure_art_director_visual_calibration_v4_ark.pdf").resolve()).replace("'", "''") + "',32);$p.Close();$a.Quit()"
    subprocess.run(["powershell", "-NoProfile", "-Command", cmd], check=True, timeout=420)
    pngs = sorted((pages / "slide").glob("*.png"), key=lambda p: int(re.search(r"(\d+)$", p.stem).group(1)))
    if len(pngs) != len(PICK): raise RuntimeError(f"COM render expected {len(PICK)} PNGs, got {len(pngs)}")
    return pngs


def contact(pngs: list[Path], path: Path, labels: list[str]) -> None:
    sample = Image.open(pngs[0]).convert("RGB"); w, gap, header = 480, 16, 28; h = round(sample.height * w / sample.width)
    sheet = Image.new("RGB", (w * 2 + gap, (h + header) * ((len(pngs) + 1) // 2)), "white"); draw = ImageDraw.Draw(sheet)
    for i, png in enumerate(pngs):
        x, y = (i % 2) * (w + gap), (i // 2) * (h + header); draw.text((x + 4, y + 4), labels[i], fill=PAD.TEXT)
        sheet.paste(Image.open(png).convert("RGB").resize((w, h)), (x, y + header))
    path.parent.mkdir(parents=True, exist_ok=True); sheet.save(path)


def compare(v4: list[Path]) -> Path:
    rows: list[tuple[str, Path]] = []
    for i, page in enumerate(PICK):
        v3_reference = page_png(V3, V3_PAGE_MAP[page]) if page in V3_PAGE_MAP else V3 / "contact_sheet.png"
        rows.extend([(f"V2 | {page:02}", page_png(BASE, page)), (f"V3 | {page:02}", v3_reference), (f"V4 Ark | {page:02}", v4[i])])
    w, h, header, gap = 360, 203, 25, 8; result = Image.new("RGB", (w * 3 + gap * 2, (h + header) * len(PICK)), "white"); draw = ImageDraw.Draw(result)
    for i, (label, path) in enumerate(rows):
        row, col = divmod(i, 3); x, y = col * (w + gap), row * (h + header); draw.text((x + 3, y + 3), label, fill=PAD.TEXT)
        result.paste(Image.open(path).convert("RGB").resize((w, h)), (x, y + header))
    target = OUT / "comparison_v2_v3_v4.png"; result.save(target); return target


def critic(direction: dict, sheet: Path, *, max_pages: int = 4) -> dict:
    target = OUT / "visual_critic.json"
    if target.is_file(): return json.loads(target.read_text(encoding="utf-8"))
    content = [{"type": "text", "text": f"You are the Visual Critic. Compare the attached V2/V3/V4 3-column sheet and real COMKING company-deck contact sheet. V4 is Ark-created. Select no more than {max_pages} V4 page numbers for a single targeted revision only if genuinely needed."}]
    content.extend(image_part(sheet, "V2 / V3 / V4 comparison: each row is an original page number"))
    content.extend(image_part(REFERENCE_SHEETS[0], "Real COMKING reference visual language"))
    content.append({"type": "text", "text": json.dumps({"global_direction": direction, "required": {"v4_better_than_v2_v3": "boolean", "selected_pages": "list of max 4 page numbers from [1,2,3,5,6,14,19,20]", "pages": "per-page critique", "revision_briefs": "page-specific defect / retain / target / locked-facts", "overall_verdict": "PASS or NEEDS_ONE_MORE_PASS"}}, ensure_ascii=False)})
    value, _ = ask_json(call_id="visual_critic", system="Return only the requested JSON Visual Critic report. Judge actual attached rendering, not intentions.", content=content, raw_file=OUT / "visual_critic_raw.json", max_tokens=6500)
    selected = [int(x) for x in value.get("selected_pages", []) if str(x).isdigit() and int(x) in PICK][:max_pages]
    value["selected_pages"] = selected; target.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8"); return value


def report(direction: dict, scenes: dict[int, dict], critic_result: dict, rendered: list[Path]) -> None:
    audits = {str(p): PAD.validate(scenes[p], set(context_for(p)["allowed_image_sources"])) for p in PICK}
    text = {str(p): text_content(scenes[p]) for p in PICK}
    manifest = {"source_v2": str((BASE / "pure_art_director_mainline_expanded_v2.pptx").relative_to(ROOT)), "source_v3": str((V3 / "pure_art_director_visual_calibration_v3.pptx").relative_to(ROOT)), "pages": list(PICK), "model": PAD.MODEL, "ark_calls": len(json.loads((OUT / "ark_calls.json").read_text(encoding="utf-8"))), "scenes": {str(p): str((OUT / "scene_graphs" / f"page_{p:02}_{NAMES[p]}.json").resolve().relative_to(OUT.resolve())) for p in PICK}, "before_text": {str(p): text_content(source_scene(p)) for p in PICK}, "after_text": text, "text_change_reason": "Ark visual hierarchy and concision only; facts and boundary conditions were locked in prompts.", "audit": audits, "critic": critic_result, "rendered_pages": [str(x.resolve().relative_to(OUT.resolve())) for x in rendered]}
    (OUT / "visual_calibration_v4_report.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Resumable Ark V4 visual calibration")
    parser.add_argument("--page", type=int, choices=PICK, help="Generate exactly one missing V4 page Scene Graph")
    parser.add_argument("--finish", action="store_true", help="Render all saved pages, critique, and make capped revisions")
    parser.add_argument("--model-max-calls", type=int, help="hard Ark-call budget, consumed before every request")
    args = parser.parse_args()
    configure_model_call_budget(args.model_max_calls)
    OUT.mkdir(parents=True, exist_ok=True)
    for path in (BASE / "contact_sheet.png", V3 / "contact_sheet.png", *REFERENCE_SHEETS):
        if not path.is_file(): raise RuntimeError(f"required visual reference missing: {path}")
    if not PAD.read_env().get("ARK_API_KEY"): raise RuntimeError("ARK_API_KEY is missing")
    say("V2/V3/reference decks inspected; Ark configuration PRESENT")
    direction = art_direction()
    if args.page:
        page_request(args.page, direction)
        return
    scenes = {page: page_request(page, direction) for page in PICK}
    if not args.finish:
        return
    say("8/8 Ark Scene Graphs saved; PowerPoint COM rendering started")
    deck = draw(scenes); rendered = render(deck); v4_sheet = OUT / "contact_sheet.png"; contact(rendered, v4_sheet, [f"V4 Ark | {p:02}" for p in PICK]); comparison_sheet = compare(rendered)
    critic_result = critic(direction, comparison_sheet)
    selected = critic_result.get("selected_pages", [])
    if selected:
        say("critic selected " + ", ".join(f"{x:02}" for x in selected) + " for one targeted revision")
        for page in selected:
            before = OUT / "revisions" / f"page_{page:02}_{NAMES[page]}_before.png"; before.parent.mkdir(parents=True, exist_ok=True); shutil.copy2(rendered[PICK.index(page)], before)
            scenes[page] = page_request(page, direction, revision=True)
        deck = draw(scenes); rendered = render(deck); contact(rendered, v4_sheet, [f"V4 Ark | {p:02}" for p in PICK]); comparison_sheet = compare(rendered)
    report(direction, scenes, critic_result, rendered)
    say("final PowerPoint COM render completed")


if __name__ == "__main__":
    try: main()
    except Exception as exc:
        print(f"V4_BLOCKED_OR_FAILED: {type(exc).__name__}: {exc}", file=sys.stderr, flush=True)
        raise
