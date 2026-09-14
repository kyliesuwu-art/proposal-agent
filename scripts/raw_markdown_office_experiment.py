"""Raw-Markdown Ark Office experiment; isolated from proposal production."""
from __future__ import annotations
import json, os, re, shutil, socket, sys, time, urllib.error, urllib.request
from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parent.parent
SOURCE=ROOT/'outputs/hospital_power_verify_20260911_attempt4_delivery/proposal.md'
OUT=ROOT/'outputs/doubao_office_test/raw_markdown'
BASE='https://ark.cn-beijing.volces.com'

PPT_SYSTEM='''你是一位顶级企业解决方案演示设计师。请基于用户提供的完整 Markdown 方案，重新组织成一份适合向企业领导汇报的专业演示文稿。你可以自由决定故事线、页数、每页主题、信息取舍、图示/表格、原始图片取舍和页面节奏。目标不是摘要原文，而是重新设计一场清晰有说服力的演示。不得编造原文不存在的项目事实和参数。输入引用只用于事实理解，最终演示不要显示引用编号、参考文献或来源列表。直接输出 PPT 专用 Markdown，不要 JSON，不要解释。以 # 表示演示标题、--- 表示换页、## 表示页标题；可保留输入已有的本地 Markdown 图片路径。'''
WORD_SYSTEM='''你是一位资深企业解决方案文档编辑与信息设计师。请基于用户提供的完整 Markdown 方案，重新整理成一份可正式交付客户的专业 Word 解决方案。允许优化标题、章节、段落、表格、信息层级、图片位置和阅读节奏。不得改变原方案事实边界或编造参数。输入引用只用于事实依据，最终文档不要显示引用编号、Sources 或 References。直接输出完整 Word Markdown，不要 JSON，不要解释。保留有效的本地 Markdown 图片路径。'''

def env():
    data={}
    for line in (ROOT/'.env').read_text(encoding='utf-8').splitlines():
        if '=' in line and not line.lstrip().startswith('#'):
            k,v=line.split('=',1);data[k.strip()]=v.strip().strip('"\'')
    return data
def post(base,key,payload,stream=False):
    req=urllib.request.Request(base.rstrip('/')+'/api/v3/chat/completions',data=json.dumps(payload,ensure_ascii=False).encode(),headers={'Authorization':'Bearer '+key,'Content-Type':'application/json'},method='POST'); start=time.perf_counter()
    try:
      with urllib.request.urlopen(req,timeout=600) as r:
       if not stream:return r.status,json.loads(r.read().decode()),time.perf_counter()-start
       parts=[]
       for raw in r:
        line=raw.decode('utf-8','replace').strip()
        if not line.startswith('data:'):continue
        value=line[5:].strip()
        if value=='[DONE]':break
        try:
         chunk=json.loads(value); delta=(chunk.get('choices')or[{}])[0].get('delta',{}).get('content','')
         if isinstance(delta,list):delta=''.join(x.get('text','') for x in delta if isinstance(x,dict))
         parts.append(delta)
        except json.JSONDecodeError:pass
       return r.status,{'choices':[{'message':{'content':''.join(parts)}}]},time.perf_counter()-start
    except (urllib.error.HTTPError,urllib.error.URLError,socket.timeout,OSError) as e:return getattr(e,'code',0),{},time.perf_counter()-start
def content(data):return str((data.get('choices')or[{}])[0].get('message',{}).get('content',''))
def clean(text):
    text=re.sub(r'\[来源:\s*[^\]]+\]','',text)
    text=re.sub(r'^图片来源[:：].*$','',text,flags=re.M)
    text=re.sub(r'^#{2,3}\s*(来源与依据|参考资料|Sources|References).*$[\s\S]*?(?=^#|\Z)','',text,flags=re.M|re.I)
    return re.sub(r'\n{3,}','\n\n',text).strip()+'\n'
def copy_assets(dst):
    src=SOURCE.parent/'assets'; target=dst/'assets'
    if target.exists():shutil.rmtree(target)
    shutil.copytree(src,target)
def slides(md):
    pieces=[p.strip() for p in re.split(r'^---\s*$',md,flags=re.M) if p.strip()]
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    for n,piece in enumerate(pieces,1):
      s=prs.slides.add_slide(prs.slide_layouts[6]); bg=s.background.fill;bg.solid();bg.fore_color.rgb=__import__('pptx').dml.color.RGBColor(246,250,252)
      lines=piece.splitlines(); title=next((x.lstrip('#').strip() for x in lines if x.startswith('#')), '项目方案')
      box=s.shapes.add_textbox(Inches(.7),Inches(.45),Inches(12),Inches(.55));tf=box.text_frame;tf.text=title;tf.paragraphs[0].runs[0].font.size=Pt(26);tf.paragraphs[0].runs[0].font.bold=True
      body_lines = [line for line in lines if not line.startswith('#') and not line.startswith('![')]
      body='\n'.join(body_lines[:2500])
      imgs=re.findall(r'!\[[^]]*\]\((assets/[^)\s]+)',piece); width=7.0 if imgs else 11.8
      b=s.shapes.add_textbox(Inches(.8),Inches(1.35),Inches(width),Inches(5.4));b.text_frame.word_wrap=True;b.text_frame.text=body
      for para in b.text_frame.paragraphs:
       for run in para.runs:run.font.size=Pt(16)
      if imgs:
       image=(md_path.parent/imgs[0]).resolve()
       if image.is_file():s.shapes.add_picture(str(image),Inches(8.15),Inches(1.45),width=Inches(4.35))
      footer=s.shapes.add_textbox(Inches(.8),Inches(7.05),Inches(11),Inches(.2));footer.text_frame.text=f'{n:02d}';footer.text_frame.paragraphs[0].alignment=PP_ALIGN.RIGHT
    return prs
def docx(md,out):
    d=Document();
    for line in md.splitlines():
      if line.startswith('#'):
       level=min(3,len(line)-len(line.lstrip('#')));d.add_heading(line.lstrip('#').strip(),level=level)
      elif line.startswith('!['):
       m=re.search(r'\((assets/[^)\s]+)',line)
       if m and (md_path.parent/m.group(1)).is_file():d.add_picture(str(md_path.parent/m.group(1)),width=Inches(5.8))
      elif line.strip() and not line.startswith('|'):d.add_paragraph(line.lstrip('- ').strip())
    d.save(out)
def run_case(name,model,thinking,stream):
    global md_path
    folder=OUT/name;folder.mkdir(parents=True,exist_ok=True);copy_assets(folder); source=SOURCE.read_text(encoding='utf-8'); key=env()['ARK_API_KEY']; base=env().get('ARK_BASE_URL',BASE).removesuffix('/api/v3')
    report={'model':model,'thinking':thinking,'stream':stream,'source_proposal_md':str(SOURCE.relative_to(ROOT)),'calls':[]}
    for label,system,file in [('slides',PPT_SYSTEM,'slides.md'),('word',WORD_SYSTEM,'word.md')]:
      payload={'model':model,'messages':[{'role':'system','content':system},{'role':'user','content':source}],'max_tokens':10000,'stream':stream}
      if thinking is not None:payload['thinking']={'type':thinking}
      code,data,secs=post(base,key,payload,stream); value=clean(content(data)); report['calls'].append({'kind':label,'http_status':code,'latency_seconds':round(secs,2),'usage':data.get('usage','NOT_REPORTED_BY_API')})
      if code!=200 or not value:raise RuntimeError(f'{name}/{label} failed HTTP {code}')
      (folder/file).write_text(value,encoding='utf-8')
    md_path=folder/'slides.md';slides((folder/'slides.md').read_text(encoding='utf-8')).save(folder/'proposal.pptx');md_path=folder/'word.md';docx((folder/'word.md').read_text(encoding='utf-8'),folder/'proposal.docx')
    (folder/'run.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
if __name__=='__main__':
  # invoked as: case model thinking|none stream
  run_case(sys.argv[1],sys.argv[2],None if sys.argv[3]=='none' else sys.argv[3],sys.argv[4]=='stream')
