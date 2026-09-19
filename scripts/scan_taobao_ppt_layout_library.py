"""Read-only scan of the private Taobao presentation-pattern asset.

No ingestion, Git staging, source mutation, or business-fact extraction.
"""
from __future__ import annotations
import json, sys
from collections import Counter
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'scripts'))
from pptx import Presentation
import ppt_template_library_v1 as lib

SOURCE=ROOT/'private_assets/ppt_templates/source/taobao/14【国企蓝色】工作汇报逻辑图.pptx'
OUT=ROOT/'outputs/ppt_layout_library_scan/taobao'

def walk(shapes):
    for s in shapes:
        yield s
        if getattr(s,'shape_type',None)==6:
            yield from walk(s.shapes)

def classify(x):
    n=x['shape_count'];t=x['text_shape_count'];p=x['picture_count'];g=x['group_count']
    tags=[]
    if p>=1 and p/max(n,1)>=.25: tags.append('IMAGE_DRIVEN')
    if p>=2: tags.append('MULTI_IMAGE')
    if g>=2 and t>=4: tags.append('CARD_OR_DIAGRAM')
    if t>=5 and p==0: tags.append('TEXT_STRUCTURE')
    if n>=12: tags.append('DENSE')
    # Structural candidates only; never infer source-template business facts.
    if g>=3 or t>=6: tags.append('MULTI_POINT')
    if 3<=t<=7 and p<=1: tags.append('CHECKLIST_OR_PROCESS')
    if p>=1 and t>=2: tags.append('IMAGE_TEXT')
    if not tags: tags=['MINIMAL']
    return tags

def capacity(x):
    if x['picture_count']>=2 and x['text_shape_count']<=4:return 'two visuals + concise narrative'
    if x['text_shape_count']>=8:return 'multi-point / dense structure'
    if x['group_count']>=3:return 'parallel cards or diagram nodes'
    if x['picture_count']>=1:return 'single visual + short explanation'
    return 'title + concise text structure'

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    prs=Presentation(SOURCE); rows=[]
    for i,slide in enumerate(prs.slides,1):
        all_shapes=list(walk(slide.shapes)); text=[]; kinds=[]
        for s in all_shapes:
            kinds.append(str(getattr(s,'shape_type',None)))
            if getattr(s,'has_text_frame',False) and s.text.strip(): text.append(s.text.strip().replace('\n',' ')[:120])
        c=Counter(kinds)
        # PICTURE=13, CHART=3, TABLE=19, GROUP=6 in python-pptx MSO enum values.
        row={'slide_number':i,'title_text':text[0] if text else '', 'shape_count':len(all_shapes),
             'text_shape_count':sum(1 for s in all_shapes if getattr(s,'has_text_frame',False)),
             'picture_count':c.get('PICTURE (13)',0),'chart_count':c.get('CHART (3)',0),
             'table_count':c.get('TABLE (19)',0),'group_count':c.get('GROUP (6)',0),
             'text_samples':text[:8]}
        row['semantic_pattern']=classify(row);row['layout_family']='TAOBAO_BLUE_LOGIC';row['content_capacity']=capacity(row)
        row['visual_density']='high' if row['shape_count']>=15 else 'medium' if row['shape_count']>=7 else 'low'
        row['has_large_image']=row['picture_count']>=1;row['has_metrics']=row['chart_count']>0
        row['has_process']='CHECKLIST_OR_PROCESS' in row['semantic_pattern'];row['has_timeline']=False
        row['has_comparison']=row['picture_count']==0 and row['text_shape_count']>=4 and row['group_count']>=2
        row['has_multi_cards']='MULTI_POINT' in row['semantic_pattern'];rows.append(row)
    patterns=Counter(tag for r in rows for tag in r['semantic_pattern'])
    raw={'source_file':'private_assets/ppt_templates/source/taobao/14【国企蓝色】工作汇报逻辑图.pptx','slide_count':len(rows),'pattern_counts':patterns,'slides':rows}
    (OUT/'slide_inventory.json').write_text(json.dumps(raw,ensure_ascii=False,indent=2,default=dict),encoding='utf-8')
    (OUT/'layout_catalog_raw.json').write_text(json.dumps({'pattern_counts':patterns,'candidate_slides':sorted(rows,key=lambda r:(r['has_multi_cards'],r['shape_count']),reverse=True)},ensure_ascii=False,indent=2,default=dict),encoding='utf-8')
    (OUT/'analysis.md').write_text('# Taobao PPT layout scan\n\n- Source is read-only private visual asset; no source business content is treated as knowledge.\n- Slides are grouped by observable layout structure, not their source titles.\n- PowerPoint COM previews are written in batches of 25 pages.\n',encoding='utf-8')
    pngs=lib.render_with_powerpoint(SOURCE,OUT/'previews/slides')
    # PowerShell's directory order is lexical (1, 10, 100); restore the
    # actual slide-number sequence before declaring contact-sheet ranges.
    pngs=sorted(pngs,key=lambda p:int(''.join(c for c in p.stem if c.isdigit())))
    for start in range(0,len(pngs),25):lib.contact_sheet(pngs[start:start+25],OUT/f'previews/contact_sheet_{start+1:03}_{min(start+25,len(pngs)):03}.png',cols=5,width=300)
if __name__=='__main__':main()
