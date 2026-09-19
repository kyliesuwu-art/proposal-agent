"""Eight-page local visual-language calibration; no model calls."""
from __future__ import annotations
import json, subprocess, re
from copy import deepcopy
from pathlib import Path
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from run_ppt_pure_art_director_editorial import ROOT, SOURCE, PAD

BASE=SOURCE/'mainline_expanded_v2'; OUT=SOURCE/'visual_calibration_v3'; PICK=[1,2,3,5,6,14,19,20]
def t(i,x,y,w,h,s,z=18,c='#18324B',b=False,a='left'):
 return {'id':i,'type':'text','x':x,'y':y,'w':w,'h':h,'text':s,'font_size':z,'font_weight':'bold' if b else 'normal','color':c,'alignment':a}
def r(i,x,y,w,h,c,z=0): return {'id':i,'type':'rectangle','x':x,'y':y,'w':w,'h':h,'fill':c,'z_order':z}
def ln(i,x1,y1,x2,y2,c='#2E75B6'): return {'id':i,'type':'line','from':[x1,y1],'to':[x2,y2],'stroke':c,'stroke_width':1.4}
def load(n): return json.loads((BASE/'scene_graphs'/(f"page_{n:02}_"+{1:'G',2:'TOC',3:'BACKGROUND',5:'H',6:'ARCH',14:'D',19:'E',20:'CLOSE'}[n]+'.json')).read_text(encoding='utf8'))
def scenes():
 s={n:load(n) for n in PICK}
 s[1]={'background':'#FFFFFF','elements':[r('glow',7.2,0,6.2,7.5,'#005EB8'),ln('energy1',7.1,1.2,12.5,5.5,'#67C6FF'),ln('energy2',8.2,6.5,12.9,2.0,'#1AA6A6'),t('eyebrow',.8,1.3,4,.3,'COMKING  智慧能源解决方案',15,'#005EB8',True),t('title',.8,2.2,6.2,1.35,'医院园区高可靠供电与智慧配电改造方案',35,'#18324B',True),t('sub',.83,3.85,5.3,.35,'医疗级高可靠智慧配电保障体系',19,'#526577'),t('tag',.83,5.85,4.5,.28,'高可靠  ·  数字化  ·  可落地',15,'#005EB8')]}
 s[2]={'background':'#F4F8FC','elements':[t('n',.78,1.05,1.4,.55,'01—04',26,'#005EB8',True),t('title',.78,1.72,5.4,.7,'汇报如何展开',35,'#18324B',True),ln('rule',.8,2.65,12.3,2.65),t('a',1.0,3.15,2.5,.35,'项目理解与建设目标',20,'#18324B',True),t('b',4.1,4.0,2.8,.35,'总体方案与关键能力',20,'#18324B',True),t('c',7.3,4.85,2.5,.35,'实施路径与保障',20,'#18324B',True),t('d',9.8,5.7,2.6,.35,'投资测算与预期价值',20,'#18324B',True),ln('p1',1.0,3.65,4.1,4.0),ln('p2',4.1,4.5,7.3,4.85),ln('p3',7.3,5.35,9.8,5.7),t('note',.8,6.55,7,.25,'从风险识别进入方案架构，再落到实施与测算边界',16,'#63798C')]}
 s[3]={'background':'#FFFFFF','elements':[r('side',.13,0,3.3,7.5,'#0D3562'),t('title',.7,1.05,2.25,1.25,'连续诊疗不能承受供电风险的滞后发现',30,'#FFFFFF',True),t('bound',.7,5.75,2.25,.6,'容量、负荷与设备台账需现场校核',15,'#A9D5F7'),t('head',4.15,1.1,7.6,.45,'项目背景与建设必要性',30,'#18324B',True),t('risk1',4.2,2.15,2.2,.3,'连续性要求',20,'#005EB8',True),t('body1',4.2,2.55,6.9,.55,'重要医疗负荷对供电连续性要求高，运行保障需要前置设计。',18),t('risk2',4.2,3.65,2.2,.3,'运行风险',20,'#D1504D',True),t('body2',4.2,4.05,6.9,.55,'设备老化、人工巡检盲区与异常发现滞后会放大运行风险。',18),t('risk3',4.2,5.15,2.5,.3,'改造方向',20,'#005EB8',True),t('body3',4.2,5.55,6.9,.55,'感知、预警与辅助决策让故障处置从事后响应向异常早期前移。',18)]}
 # Reuse technical evidence, but reduce its physical display area.
 h=s[5]; img=next(e for e in h['elements'] if e.get('type')=='image'); img.update({'x':7.0,'y':1.55,'w':5.1,'h':4.65,'crop':'contain'}); h['elements']=[img,t('title',.75,1.0,5.6,.8,'“三高”目标\n化解老旧配电风险',31,'#18324B',True),ln('blue',.78,2.75,4.9,2.75),t('hi1',.8,3.25,1.55,.3,'高可靠',21,'#005EB8',True),t('hi2',2.65,3.85,1.55,.3,'高透明',21,'#1AA6A6',True),t('hi3',4.45,4.45,1.55,.3,'高效益',21,'#D1504D',True),ln('stair1',2.1,3.55,2.65,4.15),ln('stair2',3.95,4.15,4.45,4.75),t('cap',.8,5.45,5.5,.55,'以设备升级、数字集成与专业运维形成递进的建设路径。',17,'#526577')]
 s[6]={'background':'#F5F9FC','elements':[t('title',.75,.82,11.6,.45,'总体方案架构',31,'#18324B',True),t('sub',.78,1.4,11,.3,'设备与供电基础通过感知、平台与资源调度，连接到安全运维与辅助决策',17,'#63798C'),r('base',.85,2.35,2.0,3.7,'#0D3562'),t('baseT',1.08,3.0,1.55,.8,'供电与\n设备基础',23,'#FFFFFF',True,'center'),r('sense',3.15,2.05,2.15,1.15,'#DCEEFF'),t('senseT',3.4,2.38,1.65,.35,'感知与边缘采集',18,'#005EB8',True,'center'),r('store',3.15,4.7,2.15,1.15,'#E4F7F5'),t('storeT',3.45,5.05,1.55,.35,'储能及应急资源',18,'#157D78',True,'center'),r('platform',5.85,2.75,2.4,2.05,'#005EB8'),t('platformT',6.15,3.35,1.8,.65,'智慧能源\n管理平台',23,'#FFFFFF',True,'center'),r('ops',8.85,2.05,2.9,3.8,'#FFFFFF'),t('opsT',9.15,2.55,2.3,.35,'运维、安全与决策',21,'#18324B',True,'center'),t('opsB',9.2,3.25,2.2,1.8,'在线监测\n告警处置\n辅助决策\n深化校核',17,'#526577',False,'center'),ln('d1',2.85,3.15,5.85,3.6),ln('d2',5.3,2.6,5.85,3.35),ln('d3',5.3,5.25,5.85,4.05),ln('d4',8.25,3.75,8.85,3.75)]}
 s[14]={'background':'#FFFFFF','elements':[t('title',.75,.85,11.5,.45,'三阶段实施路径',31,'#18324B',True),t('constraint',.78,1.42,11,.3,'以减少对医疗秩序的干扰为约束，组织施工窗口、联调和验收节点',17,'#63798C'),ln('route',1.2,4.1,12.0,4.1,'#005EB8'),*sum(([r(f'm{i}',x,3.78,.64,.64,c),t(f'n{i}',x,3.96,.64,.2,str(i+1).zfill(2),17,'#FFFFFF',True,'center'),t(f'h{i}',x-.45,2.45,1.65,.3,h,20,'#18324B',True,'center'),t(f'b{i}',x-.6,4.75,1.95,.65,b,16,'#526577',False,'center')] for i,(x,c,h,b) in enumerate([(1.6,'#0D3562','勘察与评审','边界校核\n施工窗口确认'),(5.4,'#005EB8','改造与部署','设备更新\n储能与平台部署'),(9.3,'#1AA6A6','联调与移交','测试验收\n培训与运维交接')])),[])]}
 s[19]={'background':'#F5F9FC','elements':[t('title',.75,.85,5.6,.45,'预期价值',31,'#18324B',True),t('action',7.0,.85,5.2,.45,'下一步工作',31,'#18324B',True),ln('split',6.63,.75,6.63,6.2),t('v1',.9,2.0,5.1,.35,'高效设施运维',23,'#005EB8',True),t('v2',.9,3.05,5.1,.35,'降本增效',23,'#157D78',True),t('v3',.9,4.1,5.1,.35,'快速响应',23,'#D1504D',True),t('note',.9,5.45,5.0,.45,'案例数字不构成本项目承诺，实际效益需独立测算。',16,'#63798C'),t('step',7.05,2.0,4.65,2.9,'01 现场勘察\n\n02 参数收集\n\n03 方案深化\n\n04 投资测算',21,'#18324B',True)]}
 s[20]={'background':'#FFFFFF','elements':[r('band',0,5.85,13.33,1.65,'#005EB8'),ln('net1',.7,6.2,4.5,7.05,'#66C6FF'),ln('net2',8.8,.8,12.7,3.2,'#1AA6A6'),t('thanks',.85,2.2,6.4,.65,'感谢聆听',42,'#18324B',True),t('more',.9,3.15,5.5,.35,'期待进一步交流',21,'#526577'),t('brand',.9,6.45,5,.3,'COMKING  智慧能源解决方案',16,'#FFFFFF',True)]}
 return s
def chrome(slide,n,src):
 slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor(*PAD.BG); clone=deepcopy(src._element); slide.shapes._spTree.insert_element_before(clone,'p:extLst'); logo=next(x for x in slide.shapes if x.name=='brand_wordmark'); logo.left,logo.top,logo.width,logo.height=Inches(.58),Inches(.17),Inches(1.55),Inches(.26); PAD.add_text(slide,t('foot',.65,7.0,4.8,.2,'医院园区智慧配电改造方案',10,'#405466')); PAD.add_text(slide,t('page',11.72,7,.9,.2,f'{n:02} / 20',10,'#1F497D',True,'right'))
def main():
 OUT.mkdir(parents=True,exist_ok=True); sc=scenes(); src=Presentation(BASE/'pure_art_director_mainline_expanded_v2.pptx'); logo=next(x for x in src.slides[0].shapes if x.name=='brand_wordmark'); prs=Presentation(); prs.slide_width=Inches(PAD.W); prs.slide_height=Inches(PAD.H); blank=prs.slide_layouts[6]
 for pos,n in enumerate(PICK,1):
  sl=prs.slides.add_slide(blank); chrome(sl,n,logo); ps={}
  for e in sorted(sc[n]['elements'],key=lambda x:x.get('z_order',10)):
   e=PAD.clamp(dict(e)); typ=e['type']; sh=PAD.add_text(sl,e) if typ=='text' else PAD.add_shape(sl,e) if typ in {'rectangle','rounded_rectangle','circle'} else PAD.add_image(sl,e) if typ=='image' else PAD.add_line(sl,e,ps); ps[e['id']]=sh
 deck=OUT/'pure_art_director_visual_calibration_v3.pptx'; prs.save(deck); pages=OUT/'pages'; subprocess.run(['powershell','-ExecutionPolicy','Bypass','-File',str(ROOT/'scripts'/'render_pptx_windows.ps1'),'-InputPath',str(deck),'-OutputDir',str(pages)],cwd=ROOT,check=True,timeout=420); cmd=f"$a=New-Object -ComObject PowerPoint.Application;$a.Visible=$true;$p=$a.Presentations.Open('{deck.resolve()}', $true,$true,$false);$p.SaveAs('{(OUT/'pure_art_director_visual_calibration_v3.pdf').resolve()}',32);$p.Close();$a.Quit()"; subprocess.run(['powershell','-NoProfile','-Command',cmd],check=True,timeout=420)
 png=sorted((pages/'slide').glob('*.png'),key=lambda x:int(re.search(r'(\d+)$',x.stem).group(1))); im=Image.open(png[0]); sheet=Image.new('RGB',(800,900),'white');
 for i,p in enumerate(png): sheet.paste(Image.open(p).convert('RGB').resize((390,219)),((i%2)*400,(i//2)*225));
 sheet.save(OUT/'contact_sheet.png'); (OUT/'visual_calibration_report.json').write_text(json.dumps({'pages':PICK,'ark_calls':0,'checked':'PowerPoint COM','low_res':'page 5 image reduced'},ensure_ascii=False,indent=2),encoding='utf8')
if __name__=='__main__': main()
