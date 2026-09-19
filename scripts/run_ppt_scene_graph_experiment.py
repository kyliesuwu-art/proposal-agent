"""Three-page Scene Graph PPT experiment.  The model owns the layout; this is a dumb executor."""
from __future__ import annotations

import base64, json, re, shutil, subprocess, sys, time, urllib.error, urllib.request
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs" / "ppt_scene_graph_experiment"
SOURCE = ROOT / "outputs" / "ppt_style_v4" / "slides.final.md"
ASSETS = ROOT / "outputs" / "ppt_style_v4" / "assets"
HYBRID_PNG = ROOT / "outputs" / "ppt_hybrid_v5" / "final" / "png" / "slide"
MODEL = "doubao-seed-2-1-pro-260628"
W, H = 13.333, 7.5
NAVY, BLUE, LIGHT_BLUE, BG, TEXT, MUTED = (31, 73, 125), (79, 129, 189), (232, 243, 248), (247, 250, 252), (24, 50, 71), (90, 107, 120)
# slides.final.md retains its cover as source item 1.  These are the exact
# existing hospital pages: platform architecture, AI platform screenshot,
# and multi-risk unattended operation respectively.
# Legacy three-page experiment mapping. Stabilization supplies its own mapping.
PAGES = {"A": 9, "B": 10, "C": 5}
STAB_PAGES = {"A": 9, "B": 10, "C": 5, "D": 12, "E": 13, "F": 14}
STAB_DIR = OUT / "stabilization"


def settings():
    values = {}
    for line in (ROOT / ".env").read_text(encoding="utf-8").splitlines():
        if "=" in line and not line.lstrip().startswith("#"):
            k, v = line.split("=", 1); values[k.strip()] = v.strip().strip("\"'")
    if not values.get("ARK_API_KEY"): raise RuntimeError("ARK_API_KEY is not present")
    return values


def ask(messages, purpose, max_tokens=16000, json_mode=False):
    """One streaming call, 360 s timeout. Logs contain only response and timing, never secrets."""
    cfg = settings(); base = cfg.get("ARK_BASE_URL", "https://ark.cn-beijing.volces.com").removesuffix("/api/v3")
    body = {"model": MODEL, "thinking": {"type": "enabled"}, "stream": True, "max_tokens": max_tokens, "messages": messages}
    if json_mode: body["response_format"]={"type":"json_object"}
    req = urllib.request.Request(base + "/api/v3/chat/completions", data=json.dumps(body, ensure_ascii=False).encode(), headers={"Authorization": "Bearer " + cfg["ARK_API_KEY"], "Content-Type": "application/json"}, method="POST")
    begun = time.perf_counter(); first = None; parts = []; event_count=0; content_chars=0; reasoning_chars=0; delta_keys=set(); finish_reasons=[]
    with urllib.request.urlopen(req, timeout=360) as response:
        for raw in response:
            line = raw.decode("utf-8", "replace").strip()
            if not line.startswith("data:"): continue
            data = line[5:].strip()
            if data == "[DONE]": break
            try:
                payload=json.loads(data); choice=(payload.get("choices") or [{}])[0]; delta_obj=choice.get("delta", {}) or {}
                event_count+=1; delta_keys.update(delta_obj.keys())
                if choice.get("finish_reason"): finish_reasons.append(choice["finish_reason"])
                thought=delta_obj.get("reasoning_content","")
                if isinstance(thought,list): thought="".join(x.get("text","") for x in thought if isinstance(x,dict))
                reasoning_chars+=len(str(thought or "")); delta=delta_obj.get("content", "")
            except json.JSONDecodeError: continue
            if isinstance(delta, list): delta = "".join(x.get("text", "") for x in delta if isinstance(x, dict))
            if delta: first = first or time.perf_counter(); parts.append(delta); content_chars+=len(str(delta))
    result = "".join(parts).strip()
    diag={"purpose":purpose,"event_count":event_count,"delta_keys":sorted(delta_keys),"content_characters":content_chars,"reasoning_characters":reasoning_chars,"finish_reasons":finish_reasons,"has_final_content":bool(result)}
    (OUT/"llm_diagnostics").mkdir(parents=True,exist_ok=True); (OUT/"llm_diagnostics"/f"{purpose}.json").write_text(json.dumps(diag,ensure_ascii=False,indent=2),encoding="utf-8")
    if not result: raise RuntimeError(f"{purpose}: empty model response; SSE={json.dumps(diag, ensure_ascii=False)}")
    meta = {"purpose": purpose, "model": MODEL, "thinking": "enabled", "stream": True, "elapsed_seconds": round(time.perf_counter()-begun, 1), "ttft_seconds": round(first-begun, 1) if first else None}
    return result, meta


def connectivity_check():
    """One redacted, minimal streaming request for checkpoint recovery."""
    cfg=settings(); base=cfg.get("ARK_BASE_URL","https://ark.cn-beijing.volces.com").removesuffix("/api/v3")
    endpoint=base+"/api/v3/chat/completions"
    body={"model":MODEL,"thinking":{"type":"enabled"},"stream":True,"max_tokens":512,"messages":[{"role":"user","content":"返回 OK"}]}
    request=urllib.request.Request(endpoint,data=json.dumps(body,ensure_ascii=False).encode(),headers={"Authorization":"Bearer "+cfg["ARK_API_KEY"],"Content-Type":"application/json"},method="POST")
    record={"request_stage":"minimal_connectivity_check","model":MODEL,"endpoint":endpoint,"thinking":"enabled","stream":True,"first_request":True}
    try:
        text,meta=ask(body["messages"],"minimal_connectivity_check",max_tokens=512)
        record.update({"http_status":200,"outcome":"success","response":text[:80],"timing":meta})
    except urllib.error.HTTPError as exc:
        raw=exc.read().decode("utf-8","replace")[:2000]
        try: payload=json.loads(raw); error=payload.get("error",payload)
        except json.JSONDecodeError: error={"message":raw}
        record.update({"http_status":exc.code,"outcome":"http_error","error_code":error.get("code") if isinstance(error,dict) else None,"error_type":error.get("type") if isinstance(error,dict) else None,"message":str(error.get("message",raw) if isinstance(error,dict) else raw)[:800]})
    except RuntimeError as exc:
        record.update({"http_status":200,"outcome":"empty_final_content","message":str(exc)[:800]})
    OUT.mkdir(parents=True,exist_ok=True)
    (OUT/"ark_connectivity_check.json").write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(record,ensure_ascii=False)); return record


def as_json(text):
    fenced = re.search(r"```(?:json)?\s*(\{[\s\S]*?\})\s*```", text, re.I)
    if fenced: return json.loads(fenced.group(1))
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end <= start: raise ValueError("model did not return JSON")
    return json.loads(text[start:end+1])


def source_briefs():
    parts = [x.strip() for x in re.split(r"^---\s*$", SOURCE.read_text(encoding="utf-8"), flags=re.M) if x.strip()]
    result = {}
    for letter, number in PAGES.items():
        piece = parts[number-1]
        title = next(x.lstrip("#").strip() for x in piece.splitlines() if x.startswith("#"))
        bullets = [re.sub(r"^\s*[*-]\s*", "", x).strip() for x in piece.splitlines() if re.match(r"^\s*[*-]\s+", x)]
        images = re.findall(r"!\[[^]]*\]\((assets/[^)\s]+)\)", piece)
        result[letter] = {"source_page": number, "title": title, "bullets": bullets, "image_candidates": images}
    return result


def scene_prompt(letter, brief):
    image_note = "无图片候选。" if not brief["image_candidates"] else "真实图片候选，只能使用这些相对路径：" + ", ".join(brief["image_candidates"])
    return f"""仅输出 JSON Scene Graph。你是医院工程汇报的 COMKING PPT Art Director，直接决定元素坐标，Python 不设计页面。
画布13.333×7.5 inch。根对象：slide_size,background,elements。12–18 个元素；每项含 id,type,x,y,w,h,z。text含 text,font_size,font_weight,color；image含 source,crop_mode；shape可含fill,line_color；arrow可用from/to坐标。不得改写事实或图片路径。
风格：白/浅灰、#1F497D/#4F81BD/#E8F3F8、正式工程汇报。含标题、COMKING、页脚和页码。标题≥24pt，正文≥12pt，主内容不进 y≥6.90 页脚区，也不覆盖右上 x≥11.55,y≤0.30 品牌区。不要卡片墙、霓虹或紫粉。
PAGE {letter}。{image_note}。事实：{json.dumps(brief,ensure_ascii=False)}
方向：A 是平台与感知/调度/分析/执行/院区对象的关系；B 让截图占50–70%并用少量文字/callout围绕；C 表现多风险感知、AI识别、平台闭环与无人值守的关系。只输出 JSON。"""


def validate(scene):
    issues = []; elements = scene.get("elements")
    if not isinstance(elements, list) or not elements: return ["elements_missing"]
    for e in elements:
        ident, typ = e.get("id"), e.get("type")
        if not ident or not typ: issues.append("element_identity_missing"); continue
        if typ in {"arrow", "line"} and "from" in e and "to" in e: continue
        try: x,y,w,h = (float(e[k]) for k in ("x","y","w","h"))
        except (KeyError, TypeError, ValueError): issues.append(f"{ident}:geometry_missing"); continue
        if w <= 0 or h <= 0 or x < -0.12 or y < -0.12 or x+w > W+.12 or y+h > H+.12: issues.append(f"{ident}:out_of_bounds")
        if typ == "text" and float(e.get("font_size", 0)) < 12: issues.append(f"{ident}:small_text")
        if typ == "image" and e.get("source") not in {"assets/image-004.jpg"}: issues.append(f"{ident}:unapproved_image")
    return issues


def clamp(e):
    if e.get("type") in {"arrow", "line"} and "from" in e: return e
    for k in ("x", "y", "w", "h"):
        if k in e: e[k] = float(e[k])
    e["w"] = min(e["w"], W); e["h"] = min(e["h"], H)
    e["x"] = max(0, min(e["x"], W-e["w"])); e["y"] = max(0, min(e["y"], H-e["h"]))
    return e


def rgb(value, default=TEXT):
    value = value or "#%02X%02X%02X" % default
    value = value.lstrip("#")
    try: return RGBColor(*tuple(int(value[n:n+2],16) for n in (0,2,4)))
    except ValueError: return RGBColor(*default)


def textbox(slide, e):
    sh = slide.shapes.add_textbox(Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"])); sh.name=str(e["id"])
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=tf.margin_right=Pt(float(e.get("margin", 0))); tf.margin_top=tf.margin_bottom=Pt(0); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment={"center":PP_ALIGN.CENTER,"right":PP_ALIGN.RIGHT}.get(e.get("align"),PP_ALIGN.LEFT)
    run=p.add_run(); run.text=str(e.get("text", "")); run.font.name="Microsoft YaHei"; run.font.size=Pt(float(e.get("font_size",16))); run.font.bold=str(e.get("font_weight","")).lower() in {"bold","700","800"}; run.font.color.rgb=rgb(e.get("color"))
    return sh


def shape(slide, e):
    kind={"rounded_rectangle":MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,"circle":MSO_AUTO_SHAPE_TYPE.OVAL,"callout":MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE}.get(e["type"],MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    sh=slide.shapes.add_shape(kind, Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"])); sh.name=str(e["id"])
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(e.get("fill"), BG)
    if e.get("line_color"):
        sh.line.color.rgb=rgb(e["line_color"]); sh.line.width=Pt(float(e.get("line_width", 0.7)))
    else: sh.line.fill.background()
    return sh


def picture(slide, e):
    path = ROOT / "outputs" / "ppt_style_v4" / e["source"]
    if not path.is_file(): raise RuntimeError(f"image unavailable: {e['source']}")
    sh=slide.shapes.add_picture(str(path), Inches(e["x"]), Inches(e["y"]), width=Inches(e["w"]), height=Inches(e["h"])); sh.name=str(e["id"])
    if e.get("crop_mode") == "cover":
        with Image.open(path) as im:
            actual=im.width/im.height; target=e["w"]/e["h"]
        if actual>target: sh.crop_left=sh.crop_right=(1-target/actual)/2
        else: sh.crop_top=sh.crop_bottom=(1-actual/target)/2
    return sh


def decorate(slide, number):
    fill=slide.background.fill; fill.solid(); fill.fore_color.rgb=RGBColor(*BG)
    bar=slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(.13), Inches(H)); bar.fill.solid(); bar.fill.fore_color.rgb=RGBColor(*BLUE); bar.line.fill.background()
    textbox(slide,{"id":"brand_wordmark","type":"text","x":11.62,"y":.08,"w":1.1,"h":.18,"text":"COMKING","font_size":9,"font_weight":"bold","color":"#1F497D","align":"right"})
    line=slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(.65), Inches(6.91), Inches(12.05), Inches(.02)); line.fill.solid(); line.fill.fore_color.rgb=RGBColor(*LIGHT_BLUE); line.line.fill.background()
    textbox(slide,{"id":"footer","type":"text","x":.65,"y":7.02,"w":3.2,"h":.18,"text":"智慧配电解决方案","font_size":9,"color":"#5A6B78"})
    textbox(slide,{"id":"page_number","type":"text","x":12.05,"y":7.00,"w":.55,"h":.2,"text":str(number).zfill(2),"font_size":10,"font_weight":"bold","color":"#4F81BD","align":"right"})


def draw_scene(scenes, deck):
    prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H); blank=prs.slide_layouts[6]
    for letter, scene in scenes.items():
        slide=prs.slides.add_slide(blank); decorate(slide, PAGES[letter]); positions={}
        ordered=sorted(scene["elements"],key=lambda e:int(e.get("z",10)))
        for raw in ordered:
            e=clamp(dict(raw)); typ=e["type"]
            if typ in {"text","number","caption"}: sh=textbox(slide,e)
            elif typ=="image": sh=picture(slide,e)
            elif typ in {"rectangle","rounded_rectangle","circle","callout"}: sh=shape(slide,e)
            elif typ in {"line","arrow"}:
                if isinstance(e.get("from"), str) and isinstance(e.get("to"), str) and e["from"] in positions and e["to"] in positions:
                    a,b=positions[e["from"]],positions[e["to"]]; x1,y1=a.left+a.width//2,a.top+a.height//2; x2,y2=b.left+b.width//2,b.top+b.height//2
                    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,x1,y1,x2,y2)
                elif (isinstance(e.get("from"), list) and len(e["from"]) == 2 and
                      isinstance(e.get("to"), list) and len(e["to"]) == 2):
                    x1,y1=Inches(float(e["from"][0])),Inches(float(e["from"][1]))
                    x2,y2=Inches(float(e["to"][0])),Inches(float(e["to"][1]))
                    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,x1,y1,x2,y2)
                else:
                    x2,y2=Inches(e["x"]+e["w"]),Inches(e["y"]+e["h"])
                    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(e["x"]), Inches(e["y"]), x2, y2)
                sh.name=str(e["id"]); sh.line.color.rgb=rgb(e.get("color"),BLUE); sh.line.width=Pt(float(e.get("line_width",1.2)))
                # python-pptx exposes no connector arrowhead API.  Use a tiny,
                # explicitly positioned native triangle as the direct executor's
                # arrow tip rather than changing the layout or routing the line.
                if typ=="arrow":
                    tip=slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE, x2-Inches(.07), y2-Inches(.07), Inches(.14), Inches(.14))
                    tip.name=str(e["id"])+"_tip"; tip.rotation=90; tip.fill.solid(); tip.fill.fore_color.rgb=rgb(e.get("color"),BLUE); tip.line.fill.background()
            else: continue
            positions[e["id"]]=sh
    prs.save(deck)


def render(deck, folder, expected_count=3):
    folder.mkdir(parents=True,exist_ok=True)
    subprocess.run(["powershell","-ExecutionPolicy","Bypass","-File",str(ROOT/"scripts"/"render_pptx_windows.ps1"),"-InputPath",str(deck),"-OutputDir",str(folder)],cwd=ROOT,check=True,timeout=360)
    image_dir=folder/"slide" if (folder/"slide").is_dir() else folder
    pngs=sorted(image_dir.glob("*.png"),key=lambda p:p.name)
    if len(pngs)!=expected_count: raise RuntimeError(f"PowerPoint COM expected {expected_count} PNGs, got {len(pngs)}")
    return pngs


def contact_sheet(pngs, output, labels=None):
    ims=[Image.open(p).convert("RGB") for p in pngs]; thumb_w=640; gap=20; header=38; thumb_h=round(ims[0].height*thumb_w/ims[0].width)
    sheet=Image.new("RGB",(thumb_w*len(ims)+gap*(len(ims)-1),thumb_h+header),(255,255,255)); draw=ImageDraw.Draw(sheet)
    for i,im in enumerate(ims):
        x=i*(thumb_w+gap); draw.text((x+6,8),(labels or [p.stem for p in pngs])[i],fill=(20,50,71)); sheet.paste(im.resize((thumb_w,thumb_h)),(x,header))
    sheet.save(output)


def image_data(path): return "data:image/png;base64,"+base64.b64encode(path.read_bytes()).decode()


def review_prompt(pngs):
    content=[{"type":"text","text":"你是严谨的医院工程汇报PPT视觉审稿人。仅评估实际PowerPoint渲染图：信息层级、平衡、图片占比、元素位置、留白、视觉重心、专业度、是否仍有卡片模板感。每页只提出能通过一次坐标/尺寸/层级调整解决的问题。不要改事实和文字。严格输出JSON：{pages:[{letter,findings:[...],revision_instructions:[...]}]}。"}]
    for p in pngs: content.append({"type":"image_url","image_url":{"url":image_data(p)}})
    return [{"role":"user","content":content}]


def revise_scene(scene, review, letter):
    prompt=f"""你是原页面设计师。根据真实渲染后的审稿意见，仅调整以下 Scene Graph 的坐标、尺寸、z、图文比例、留白、图形关系。不得修改任何 text 字符串、不得添加图片来源、不得删除品牌页脚。严格输出完整 JSON Scene Graph，不要说明。\n审稿意见：{json.dumps(review,ensure_ascii=False)}\n原SceneGraph：{json.dumps(scene,ensure_ascii=False)}"""
    text,meta=ask([{"role":"user","content":prompt}],f"revision_{letter}",json_mode=True)
    return as_json(text),meta


def comparison(final_pngs):
    hybrid=[]
    for n in (9,10,5):
        found=next((p for p in HYBRID_PNG.glob("*.png") if re.search(rf"{n}$",p.stem)),None)
        if not found: raise RuntimeError(f"Hybrid V5 PNG for page {n} unavailable")
        hybrid.append(found)
    rows=[]
    for old,new,letter in zip(hybrid,final_pngs,("PAGE A","PAGE B","PAGE C")):
        rows.extend([(old,"Hybrid V5 · "+letter),(new,"Scene Graph Final · "+letter)])
    ims=[Image.open(p).convert("RGB") for p,_ in rows]; tw=620; th=round(ims[0].height*tw/ims[0].width); sheet=Image.new("RGB",(tw*2+24,(th+34)*3),(255,255,255)); draw=ImageDraw.Draw(sheet)
    for i,(im,(_,label)) in enumerate(zip(ims,rows)):
        r,c=divmod(i,2); x=c*(tw+24); y=r*(th+34); draw.text((x+5,y+6),label,fill=(20,50,71)); sheet.paste(im.resize((tw,th)),(x,y+28))
    output=OUT/"final"/"png"/"comparison_vs_hybrid_v5.png"; sheet.save(output); return output


def stabilization_locks():
    """Build content locks from the actual source slide, never from a slide index guess."""
    pieces=[x.strip() for x in re.split(r"^---\s*$",SOURCE.read_text(encoding="utf-8"),flags=re.M) if x.strip()]
    distinct={
        "A":["智慧能源管理平台","数字孪生","AI调度","调度监控","计量计费","节能分析","虚拟电厂"],
        "B":["AI预训练运维模型","自然语言交互","告警信息","根因","边缘计算","移动运维"],
        "C":["过压","欠压","温升","过载","漏电","AI摄像头","烟火","无人值守"],
        "D":["第一阶段","第二阶段","第三阶段","勘察","评审","施工","验收","移交","代维"],
        "E":["可靠性","运维提效","节能降碳","40%","30%","20%","8%-12%","10%","行业案例参考"],
        "F":["负荷边界","并网条件","空间资源","市场规则","管理要求","总供电容量","峰值负荷","电价","预算"],
    }
    locks={}
    for page_id,source_page in STAB_PAGES.items():
        piece=pieces[source_page-1]; title=next(x.lstrip("#").strip() for x in piece.splitlines() if x.startswith("#"))
        bullets=[re.sub(r"^\s*[*-]\s*","",x).strip() for x in piece.splitlines() if re.match(r"^\s*[*-]\s+",x)]
        images=re.findall(r"!\[[^]]*\]\((assets/[^)\s]+)\)",piece)
        forbidden=sorted({term for other,terms in distinct.items() if other!=page_id for term in terms})
        locks[page_id]={"page_id":page_id,"source_page":source_page,"page_title":title,
            "primary_message":bullets[0] if bullets else title,"allowed_text_blocks":bullets,
            "allowed_facts":bullets,"allowed_images":images,"must_preserve_terms":distinct[page_id],
            "forbidden_cross_page_terms":forbidden,"semantic_page_id":source_page}
    return locks


def stabilization_prompt(lock):
    return f"""只输出 JSON Scene Graph。你是 COMKING 医院工程汇报的 Art Director，直接定义画布中每个元素的位置、尺寸、层级和关系，Python 只执行。
画布 13.333×7.5 inch。根对象含 slide_size, background, elements。元素可用 text/image/rectangle/rounded_rectangle/circle/line/arrow/callout/group；每项有 id,type,x,y,w,h,z；text 另有 text,font_size,font_weight,color；image 只能使用 allowed_images。正式蓝白工程风格。渲染器固定添加 COMKING、页脚和页码，Scene Graph 不得重复它们。标题≥24pt，正文≥12pt；主内容不得进入 y>=6.90 页脚区。
BUSINESS CONTENT IS LOCKED：只能基于下面 page_content_lock 压缩或改写展示语言；不得引入任何其他页面业务内容、事实、数字或图片。不要解释，不要模板卡片墙。\n{json.dumps(lock,ensure_ascii=False)}"""


def semantic_validate(scene, lock):
    """Deterministic cross-page guard: source facts/numbers cannot drift between pages."""
    text="\n".join(str(e.get("text","")) for e in scene.get("elements",[]) if e.get("type") in {"text","caption","number"})
    issues=[]
    allowed_source=" ".join([lock["page_title"],*lock["allowed_text_blocks"],*lock["allowed_facts"]])
    for term in lock["forbidden_cross_page_terms"]:
        # A term shared with the current source content is not cross-page
        # pollution. The lock, not a globally unique vocabulary, is authority.
        if term and term not in allowed_source and term in text: issues.append(f"cross_page_term:{term}")
    allowed_numbers=set(re.findall(r"\d+(?:\.\d+)?%?|\d+kW|\d+kWh|\d+万", " ".join([lock["page_title"],*lock["allowed_facts"]])))
    for number in re.findall(r"\d+(?:\.\d+)?%?|\d+kW|\d+kWh|\d+万",text):
        if number not in allowed_numbers and number not in {"13.333","7.5","1","2","3","4","5","6"}: issues.append(f"unlocked_number:{number}")
    if not any(term in text for term in lock["must_preserve_terms"]): issues.append("missing_page_identity_term")
    return sorted(set(issues))


def structural_validate(scene, lock):
    issues=[]; seen=set(); allowed=set(lock["allowed_images"])
    for e in scene.get("elements",[]):
        ident=e.get("id"); typ=e.get("type")
        if ident in seen: issues.append(f"duplicate_id:{ident}")
        seen.add(ident)
        if typ in {"arrow","line"} and isinstance(e.get("from"),list) and isinstance(e.get("to"),list): continue
        try: x,y,w,h=(float(e[k]) for k in ("x","y","w","h"))
        except Exception: issues.append(f"geometry:{ident}"); continue
        if min(w,h)<=0 or x<-.12 or y<-.12 or x+w>W+.12 or y+h>H+.12: issues.append(f"bounds:{ident}")
        if typ=="text" and float(e.get("font_size",0))<12: issues.append(f"small_text:{ident}")
        if typ=="image" and e.get("source") not in allowed: issues.append(f"image:{ident}")
        if y+h>6.90 and typ not in {"line"} and str(ident) not in {"footer","page_number"}: issues.append(f"footer_collision:{ident}")
    return issues


def normalize_scene(scene):
    """Schema compatibility only; never supplies a layout or changes content."""
    normalized=json.loads(json.dumps(scene))
    for e in normalized.get("elements",[]):
        if e.get("type")=="image" and "source" not in e and "src" in e: e["source"]=e["src"]
        if e.get("type") in {"rectangle","rounded_rectangle","circle","callout"} and "fill" not in e and "color" in e: e["fill"]=e["color"]
        if "align" not in e and "text_align" in e: e["align"]=e["text_align"]
        if e.get("type")=="image" and "crop_mode" not in e: e["crop_mode"]="contain"
    return normalized


def draw_stabilization(scenes, deck):
    global PAGES
    old=PAGES; PAGES={k:STAB_PAGES[k] for k in scenes}
    try: draw_scene(scenes,deck)
    finally: PAGES=old


def stabilization_review(pngs, locks):
    content=[{"type":"text","text":"你是医院工程汇报PPT视觉审稿人。BUSINESS CONTENT IS LOCKED，不得建议改变文字、事实、数字或业务含义。逐页仅提出坐标、尺寸、字号、对齐、层级、裁切、箭头、留白、视觉强调的修改。严格输出JSON：{pages:[{page_id,findings,revision_instructions}]}。"}]
    for p in pngs: content.append({"type":"image_url","image_url":{"url":image_data(p)}})
    return content


def stabilize_revision(scene, lock, review):
    prompt=f"""仅输出完整 JSON Scene Graph。BUSINESS CONTENT IS LOCKED。text 字段必须逐字符保持原 Scene Graph 不变；不得增加/删除任何业务事实或数字，不得改变图片来源。只允许 x,y,w,h,font_size,align,z,crop_mode、arrow from/to、fill/stroke/opacity 等视觉属性变化。\n内容锁：{json.dumps(lock,ensure_ascii=False)}\n审查：{json.dumps(review,ensure_ascii=False)}\n原图：{json.dumps(scene,ensure_ascii=False)}"""
    text,meta=ask([{"role":"user","content":prompt}],f"stabilization_revision_{lock['page_id']}",json_mode=True)
    return as_json(text),meta


def stabilization_comparison(final_pngs, locks):
    rows=[]
    for page_id,new in zip(locks,final_pngs):
        source=locks[page_id]["semantic_page_id"]
        old=next((p for p in HYBRID_PNG.glob("*.png") if re.search(rf"{source}$",p.stem)),None)
        rows.append((old,new,page_id,source))
    tw=580; gap=22; header=32; opened=[]
    for old,new,page_id,source in rows:
        left=Image.open(old).convert("RGB") if old else Image.new("RGB",(1600,900),(238,238,238)); right=Image.open(new).convert("RGB")
        opened.append((left,right,page_id,source))
    th=round(opened[0][0].height*tw/opened[0][0].width); sheet=Image.new("RGB",(tw*2+gap,(th+header)*len(opened)),"white"); draw=ImageDraw.Draw(sheet)
    for row,(left,right,page_id,source) in enumerate(opened):
        y=row*(th+header); draw.text((4,y+5),f"Hybrid V5 · source {source}" if rows[row][0] else "NO_EXACT_BASELINE",fill=(20,50,71)); draw.text((tw+gap+4,y+5),f"Scene Graph · {page_id} · source {source}",fill=(20,50,71)); sheet.paste(left.resize((tw,th)),(0,y+header)); sheet.paste(right.resize((tw,th)),(tw+gap,y+header))
    path=STAB_DIR/"final"/"comparison_vs_hybrid_v5.png"; sheet.save(path); return path


def stabilization_main():
    for name in ("content_locks","plans","raw","first_pass","review","final"): (STAB_DIR/name).mkdir(parents=True,exist_ok=True)
    locks=stabilization_locks(); scenes={}; report={"pages":{},"model":MODEL,"thinking":"enabled","stream":True,"stage":"generating"}
    for page_id,lock in locks.items():
        (STAB_DIR/"content_locks"/f"page_{page_id}_content_lock.json").write_text(json.dumps(lock,ensure_ascii=False,indent=2),encoding="utf-8")
        raw_path=STAB_DIR/"raw"/f"page_{page_id}_scene_raw.json"; plan_path=STAB_DIR/"plans"/f"page_{page_id}_scene.json"
        if raw_path.exists(): scene=json.loads(raw_path.read_text(encoding="utf-8")); meta={"resumed_raw":True}
        else:
            text,meta=ask([{"role":"user","content":stabilization_prompt(lock)}],f"stabilization_scene_{page_id}",json_mode=True)
            scene=as_json(text); raw_path.write_text(json.dumps(scene,ensure_ascii=False,indent=2),encoding="utf-8")
        scene=normalize_scene(scene)
        issues=structural_validate(scene,lock)+semantic_validate(scene,lock)
        if issues: raise RuntimeError(f"{page_id} raw validation failed: {issues}")
        plan_path.write_text(json.dumps(scene,ensure_ascii=False,indent=2),encoding="utf-8"); scenes[page_id]=scene; report["pages"][page_id]={"source_page":lock["source_page"],"elements":len(scene.get("elements",[])),"raw_validation":issues,"generation":meta}
    report["stage"]="first_pass"; first=STAB_DIR/"first_pass"/"scene_graph_stabilization_first_pass.pptx"; draw_stabilization(scenes,first); pngs=render(first,STAB_DIR/"first_pass"/"png",expected_count=len(scenes))
    review_text,review_meta=ask([{"role":"user","content":stabilization_review(pngs,locks)}],"stabilization_visual_review",max_tokens=6000,json_mode=True); review=as_json(review_text); (STAB_DIR/"review"/"visual_review.json").write_text(json.dumps({"review":review,"meta":review_meta},ensure_ascii=False,indent=2),encoding="utf-8")
    revised={}
    for page_id,scene in scenes.items():
        item=next((x for x in review.get("pages",[]) if x.get("page_id")==page_id),{"page_id":page_id,"findings":[],"revision_instructions":[]})
        updated,meta=stabilize_revision(scene,locks[page_id],item)
        if [e.get("text") for e in updated.get("elements",[]) if "text" in e] != [e.get("text") for e in scene.get("elements",[]) if "text" in e]: raise RuntimeError(f"{page_id} revision changed locked text")
        issues=structural_validate(updated,locks[page_id])+semantic_validate(updated,locks[page_id])
        if issues: raise RuntimeError(f"{page_id} revision validation failed: {issues}")
        (STAB_DIR/"plans"/f"page_{page_id}_scene_final.json").write_text(json.dumps(updated,ensure_ascii=False,indent=2),encoding="utf-8"); revised[page_id]=updated; report["pages"][page_id].update({"final_elements":len(updated.get("elements",[])),"revision":meta,"review_findings":item.get("findings",[]),"final_validation":issues})
    final=STAB_DIR/"final"/"scene_graph_stabilization.pptx"; draw_stabilization(revised,final); final_pngs=render(final,STAB_DIR/"final"/"png",expected_count=len(revised)); contact_sheet(final_pngs,STAB_DIR/"final"/"scene_graph_stabilization_contact_sheet.png",[f"PAGE {p} · source {locks[p]['source_page']}" for p in revised]); compare=stabilization_comparison(final_pngs,locks)
    report.update({"stage":"complete","com_pngs":len(final_pngs),"semantic_mapping":{p:locks[p]["semantic_page_id"] for p in locks},"final_pptx":str(final.relative_to(ROOT)),"contact_sheet":str((STAB_DIR/"final"/"scene_graph_stabilization_contact_sheet.png").relative_to(ROOT)),"comparison":str(compare.relative_to(ROOT))})
    (STAB_DIR/"semantic_validation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8"); (STAB_DIR/"scene_graph_validation_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    (STAB_DIR/"experiment_report.md").write_text("# Scene Graph Stabilization\n\n- Six semantic-locked hospital pages.\n- PowerPoint COM rendered all pages.\n",encoding="utf-8")
    print(json.dumps({"status":"complete","pptx":str(final),"pages":len(revised)},ensure_ascii=False))


def main():
    for name in ("source","plans","first_pass","review","final"): (OUT/name).mkdir(parents=True,exist_ok=True)
    briefs=source_briefs(); (OUT/"source"/"content_briefs.json").write_text(json.dumps(briefs,ensure_ascii=False,indent=2),encoding="utf-8")
    scenes={}; validation={"model":MODEL,"thinking":"enabled","pages":{},"stage":"initial"}
    for letter,brief in briefs.items():
        plan_path=OUT/"plans"/f"page_{letter}_scene.json"
        if plan_path.exists():
            scene=json.loads(plan_path.read_text(encoding="utf-8")); meta={"purpose":f"scene_{letter}","reused_valid_saved_scene":True}
        else:
            text,meta=ask([{"role":"user","content":scene_prompt(letter,brief)}],f"scene_{letter}",json_mode=True)
            try: scene=as_json(text)
            except (ValueError, json.JSONDecodeError):
                # Serialization repair only: preserve the page design verbatim.
                repair=("将下列内容逐字符修复为合法 JSON。不得改变设计、文字、坐标、颜色或元素，"
                        "只修复 JSON 转义、逗号或引号。只输出 JSON 对象。\n"+text)
                repaired,repair_meta=ask([{"role":"user","content":repair}],f"scene_{letter}_json_repair",json_mode=True)
                scene=as_json(repaired); meta["serialization_repair_call"]=repair_meta
        # Save a model response before safety validation.  If a subsequent
        # model-only correction is rate-limited, a retry resumes from exactly
        # this design rather than creating an alternative page.
        plan_path.write_text(json.dumps(scene,ensure_ascii=False,indent=2),encoding="utf-8")
        problems=validate(scene)
        if problems:
            # Safety correction is a constrained model revision, never a Python
            # re-layout: the author only moves/resizes the cited elements.
            safety=("你的 Scene Graph 未通过硬性执行边界检查。仅调整有问题元素的 x,y,w,h 或 font_size，"
                    "不得改变任何 text、图片、视觉结构或其他元素。画布13.333x7.5；正文>=12pt；"
                    "主内容不得占用页脚 y>=6.90。严格输出完整合法JSON。\n问题："+json.dumps(problems,ensure_ascii=False)+"\n原SceneGraph："+json.dumps(scene,ensure_ascii=False))
            corrected,correct_meta=ask([{"role":"user","content":safety}],f"scene_{letter}_safety_revision",json_mode=True)
            scene=as_json(corrected); meta["safety_revision_call"]=correct_meta; problems=validate(scene)
        if problems: raise RuntimeError(f"scene {letter} validation requires model revision: {problems}")
        scenes[letter]=scene; plan_path.write_text(json.dumps(scene,ensure_ascii=False,indent=2),encoding="utf-8"); validation["pages"][letter]={"initial_elements":len(scene["elements"]),"initial_validation":problems,"initial_call":meta}
    first=OUT/"first_pass"/"scene_graph_first_pass.pptx"; draw_scene(scenes,first); initial_pngs=render(first,OUT/"first_pass"/"png")
    review_text,review_meta=ask(review_prompt(initial_pngs),"visual_review",max_tokens=4000); review=as_json(review_text); (OUT/"review"/"visual_review.json").write_text(json.dumps({"review":review,"call":review_meta},ensure_ascii=False,indent=2),encoding="utf-8")
    revised={}
    for letter,scene in scenes.items():
        item=next((x for x in review.get("pages",[]) if x.get("letter")==letter),{"letter":letter,"findings":[],"revision_instructions":[]})
        updated,meta=revise_scene(scene,item,letter); problems=validate(updated)
        if problems: raise RuntimeError(f"revised scene {letter} invalid: {problems}")
        revised[letter]=updated; (OUT/"plans"/f"page_{letter}_scene_final.json").write_text(json.dumps(updated,ensure_ascii=False,indent=2),encoding="utf-8"); validation["pages"][letter].update({"final_elements":len(updated["elements"]),"final_validation":problems,"revision_call":meta,"review_findings":item.get("findings",[])})
    final=OUT/"final"/"scene_graph_final.pptx"; draw_scene(revised,final); final_pngs=render(final,OUT/"final"/"png"); contact_sheet(final_pngs,OUT/"final"/"png"/"contact_sheet.png",["PAGE A · Overall architecture","PAGE B · AI operations platform","PAGE C · Unattended safety"]); compare=comparison(final_pngs)
    validation.update({"stage":"complete","com_rendered_pngs":len(final_pngs),"final_pptx":str(final.relative_to(ROOT)),"comparison":str(compare.relative_to(ROOT))}); (OUT/"scene_graph_validation.json").write_text(json.dumps(validation,ensure_ascii=False,indent=2),encoding="utf-8")
    report=["# PPT Scene Graph Experiment", "", "- Scope: three existing hospital pages; no RAG, proposal, database, V5 or template-library modification.", "- Model: doubao-seed-2-1-pro-260628, thinking enabled, streaming.", "- Renderer: direct Scene Graph executor with no page-type layout dispatch.", "- PowerPoint COM: 3/3 final PNGs.", "", "## Pages", ""]
    report += [f"- {letter}: source page {briefs[letter]['source_page']} — {briefs[letter]['title']}" for letter in "ABC"]
    report += ["", "## Deliverables", "", f"- Final deck: `{final.relative_to(ROOT)}`", f"- Contact sheet: `outputs/ppt_scene_graph_experiment/final/png/contact_sheet.png`", f"- Comparison: `{compare.relative_to(ROOT)}`"]
    (OUT/"experiment_report.md").write_text("\n".join(report)+"\n",encoding="utf-8")
    print(json.dumps({"status":"complete","final":str(final),"pngs":len(final_pngs)},ensure_ascii=False))

if __name__=="__main__":
    try:
        if "--connectivity-check" in sys.argv: connectivity_check()
        elif "--stabilization" in sys.argv: stabilization_main()
        else: main()
    except Exception as exc:
        # Preserve a redacted, actionable experiment failure without leaking
        # request payloads or credentials; useful when an external SSE stream
        # is interrupted before a plan can be saved.
        OUT.mkdir(parents=True, exist_ok=True)
        (OUT / "failure.json").write_text(json.dumps({"stage": "failed", "exception": type(exc).__name__, "message": str(exc)[:500]}, ensure_ascii=False, indent=2), encoding="utf-8")
        raise
