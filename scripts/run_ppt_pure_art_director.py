"""Pure Art Director PPT experiment.

This deliberately contains no page-type classifier, template, slot map, or
layout repair.  Seed receives only content, permitted real assets, a visual
goal, and the COMKING brand context; PowerPoint is a literal Scene Graph
executor with minimal hard-error checks.
"""
from __future__ import annotations

import argparse
import base64
import json
import os
import re
import shutil
import subprocess
import time
import urllib.request
from collections import Counter
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "outputs" / "ppt_pure_art_director"
SOURCE = ROOT / "outputs" / "ppt_style_v4" / "slides.final.md"
ASSET_ROOT = ROOT / "outputs" / "ppt_style_v4"
HYBRID_PNG = ROOT / "outputs" / "ppt_hybrid_v5" / "final" / "png" / "slide"
STAB_ROOT = ROOT / "outputs" / "ppt_scene_graph_experiment"
MODEL = "doubao-seed-2-1-pro-260628"
W, H = 13.333, 7.5
BG = (247, 250, 252)
BLUE = (79, 129, 189)
NAVY = (31, 73, 125)
LIGHT_BLUE = (232, 243, 248)
TEXT = (24, 50, 71)

# The source page numbers are intentionally explicit evidence mapping, not a
# layout decision.  They are the same semantic hospital subjects used for the
# requested comparison.
PAGES = {"A": 9, "B": 10, "C": 5, "D": 12, "E": 13, "F": 14}


def read_env() -> dict[str, str]:
    """Read local defaults, then inherit the explicitly loaded runtime environment."""
    values: dict[str, str] = {}
    env_file = ROOT / ".env"
    if env_file.is_file():
        for line in env_file.read_text(encoding="utf-8").splitlines():
            if "=" in line and not line.lstrip().startswith("#"):
                key, value = line.split("=", 1)
                values[key.strip()] = value.strip().strip("\"'")
    for key in ("ARK_API_KEY", "ARK_BASE_URL"):
        if os.environ.get(key):
            values[key] = os.environ[key]
    if not values.get("ARK_API_KEY"):
        raise RuntimeError("ARK_API_KEY is not present in the runtime environment or .env")
    return values


def ask(messages: list[dict], purpose: str, *, max_tokens: int = 12000) -> tuple[str, dict]:
    """Make one complete JSON request without logging credentials.

    Every PPT consumer waits for a complete JSON object before it can proceed.
    A non-stream response removes an unnecessary SSE read boundary without
    changing the model, prompt, timeout, or JSON-object contract.
    """
    cfg = read_env()
    base = cfg.get("ARK_BASE_URL", "https://ark.cn-beijing.volces.com").removesuffix("/api/v3")
    payload = {
        "model": MODEL,
        "thinking": {"type": "enabled"},
        "stream": False,
        "max_tokens": max_tokens,
        "response_format": {"type": "json_object"},
        "messages": messages,
    }
    request = urllib.request.Request(
        base + "/api/v3/chat/completions",
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={"Authorization": "Bearer " + cfg["ARK_API_KEY"], "Content-Type": "application/json"},
        method="POST",
    )
    started = time.perf_counter()
    with urllib.request.urlopen(request, timeout=360) as response:
        body = json.loads(response.read().decode("utf-8"))
        choice = (body.get("choices") or [{}])[0]
        result = (choice.get("message") or {}).get("content", "")
    if isinstance(result, list):
        result = "".join(item.get("text", "") for item in result if isinstance(item, dict))
    result = str(result).strip()
    if not result:
        raise RuntimeError(f"{purpose}: model returned no final content")
    return result, {
        "purpose": purpose,
        "model": MODEL,
        "thinking": "enabled",
        "stream": False,
        "timeout_seconds": 360,
        "elapsed_seconds": round(time.perf_counter() - started, 1),
        "ttft_seconds": None,
        "sse_events": 0,
    }


def as_json(raw: str) -> dict:
    fenced = re.search(r"```(?:json)?\s*(\{[\s\S]*?\})\s*```", raw, re.I)
    payload = fenced.group(1) if fenced else raw[raw.find("{") : raw.rfind("}") + 1]
    if not payload:
        raise ValueError("model did not return a JSON object")
    value = json.loads(payload)
    if not isinstance(value, dict):
        raise ValueError("model JSON root is not an object")
    return value


def source_briefs() -> dict[str, dict]:
    pieces = [x.strip() for x in re.split(r"^---\s*$", SOURCE.read_text(encoding="utf-8"), flags=re.M) if x.strip()]
    result: dict[str, dict] = {}
    for letter, page_no in PAGES.items():
        piece = pieces[page_no - 1]
        title = next(x.lstrip("#").strip() for x in piece.splitlines() if x.startswith("#"))
        bullets = [re.sub(r"^\s*[*-]\s*", "", x).strip() for x in piece.splitlines() if re.match(r"^\s*[*-]\s+", x)]
        paths = re.findall(r"!\[[^]]*\]\((assets/[^)\s]+)\)", piece)
        assets = []
        for source in paths:
            if not (ASSET_ROOT / source).is_file():
                raise RuntimeError(f"source image does not exist: {source}")
            assets.append({
                "image_path": source,
                "description": "真实医院方案来源图片，内容以图片文件本身为准",
                "required": letter == "B",
            })
        if letter == "B" and not any(x["image_path"] == "assets/image-004.jpg" for x in assets):
            raise RuntimeError("Page B must retain assets/image-004.jpg")
        result[letter] = {
            "source_page": page_no,
            "title": title,
            "primary_message": bullets[0] if bullets else title,
            "facts_and_bullets": bullets,
            "image_assets": assets,
        }
    return result


def first_pass_prompt(brief: dict) -> list[dict]:
    # The user payload deliberately has exactly the four requested categories.
    user = {
        "PAGE CONTENT": {
            "page_title": brief["title"],
            "primary_message": brief["primary_message"],
            "necessary_facts_and_bullets": brief["facts_and_bullets"],
        },
        "IMAGE ASSETS": brief["image_assets"] or "本页没有可用的真实图片，不要编造或引用其他图片。",
        "VISUAL GOAL": "专业、简洁，像大型能源企业给医院管理层做正式工程汇报。视觉重点清晰，不要像自动生成的卡片拼贴；图片应真正参与设计；信息层级明确，留白合理。",
        "BRAND": "COMKING；品牌蓝；浅色背景；正式工程汇报；统一 Logo 与 footer。",
    }
    system = """你是 PPT Art Director。请自行完成整页视觉设计并只输出完整、合法的 JSON Scene Graph，不要解释。画布为 16:9 PowerPoint，slide width=13.333、slide height=7.5（单位 inch）。根对象必须包含 background 和 elements。每个 element 至少包含 id、type、x、y、w、h；按需可包含 text、font_size、font_weight、color、fill、stroke、alignment、image_source、crop、z_order、from、to、children。可使用 text、image、rectangle、rounded_rectangle、circle、line、arrow、callout、group 等元素类型。只使用 IMAGE ASSETS 中的 image_path，且 required=true 的图片必须使用。"""
    return [{"role": "system", "content": system}, {"role": "user", "content": json.dumps(user, ensure_ascii=False)}]


def validate(scene: dict, allowed_images: set[str]) -> dict:
    issues: list[str] = []
    elements = scene.get("elements")
    if not isinstance(elements, list) or not elements:
        return {"hard_errors": ["elements_missing"], "preclamp_out_of_bounds": [], "small_font": [], "duplicate_ids": [], "overlap_pairs": []}
    ids: list[str] = []
    preclamp: list[str] = []
    tiny: list[str] = []
    boxes: list[tuple[str, float, float, float, float]] = []
    for element in elements:
        if not isinstance(element, dict):
            issues.append("element_not_object")
            continue
        ident, typ = str(element.get("id", "")), element.get("type")
        if not ident or not typ:
            issues.append("element_identity_missing")
            continue
        ids.append(ident)
        if typ in {"line", "arrow"} and isinstance(element.get("from"), (list, str)) and isinstance(element.get("to"), (list, str)):
            continue
        try:
            x, y, w, h = (float(element[k]) for k in ("x", "y", "w", "h"))
        except (KeyError, TypeError, ValueError):
            issues.append(f"{ident}:geometry_missing")
            continue
        if w <= 0 or h <= 0:
            issues.append(f"{ident}:nonpositive_geometry")
        if x < -0.15 or y < -0.15 or x + w > W + 0.15 or y + h > H + 0.15:
            preclamp.append(ident)
        if typ == "text":
            try:
                if float(element.get("font_size", 0)) < 12:
                    tiny.append(ident)
            except (TypeError, ValueError):
                issues.append(f"{ident}:font_size_invalid")
        if typ == "image" and element.get("image_source") not in allowed_images:
            issues.append(f"{ident}:unapproved_or_missing_image")
        if typ not in {"line", "arrow"}:
            boxes.append((ident, x, y, w, h))
    duplicates = sorted(k for k, v in Counter(ids).items() if v > 1)
    issues.extend(f"duplicate_id:{x}" for x in duplicates)
    known = set(ids)
    for element in elements:
        if element.get("type") == "arrow" and isinstance(element.get("from"), str) and element["from"] not in known:
            issues.append(f"{element.get('id')}:arrow_from_missing")
        if element.get("type") == "arrow" and isinstance(element.get("to"), str) and element["to"] not in known:
            issues.append(f"{element.get('id')}:arrow_to_missing")
    overlaps = []
    for i, (aid, ax, ay, aw, ah) in enumerate(boxes):
        for bid, bx, by, bw, bh in boxes[i + 1:]:
            if min(ax + aw, bx + bw) - max(ax, bx) > 0.25 and min(ay + ah, by + bh) - max(ay, by) > 0.18:
                overlaps.append(f"{aid}~{bid}")
    return {"hard_errors": issues, "preclamp_out_of_bounds": preclamp, "small_font": tiny, "duplicate_ids": duplicates, "overlap_pairs": overlaps}


def clamp(element: dict) -> dict:
    """Permit minor boundary correction only; it never repositions for aesthetics."""
    if element.get("type") in {"line", "arrow"} and "from" in element and "to" in element:
        return element
    for key in ("x", "y", "w", "h"):
        element[key] = float(element[key])
    element["w"] = min(element["w"], W)
    element["h"] = min(element["h"], H)
    element["x"] = max(0, min(element["x"], W - element["w"]))
    element["y"] = max(0, min(element["y"], H - element["h"]))
    return element


def rgb(value, default=TEXT):
    raw = str(value or "#%02X%02X%02X" % default).lstrip("#")
    try:
        return RGBColor(*(int(raw[i:i + 2], 16) for i in (0, 2, 4)))
    except ValueError:
        return RGBColor(*default)


def add_text(slide, e: dict):
    shape = slide.shapes.add_textbox(Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
    shape.name = str(e["id"])
    frame = shape.text_frame
    frame.clear(); frame.word_wrap = True
    frame.margin_left = frame.margin_right = Pt(float(e.get("margin", 0)))
    frame.margin_top = frame.margin_bottom = Pt(0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(e.get("alignment"), PP_ALIGN.LEFT)
    run = paragraph.add_run(); run.text = str(e.get("text", ""))
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(float(e.get("font_size", 16)))
    run.font.bold = str(e.get("font_weight", "")).lower() in {"bold", "700", "800", "900"}
    run.font.color.rgb = rgb(e.get("color"))
    return shape


def add_shape(slide, e: dict):
    kind = {"rounded_rectangle": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, "circle": MSO_AUTO_SHAPE_TYPE.OVAL, "callout": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE}.get(e["type"], MSO_AUTO_SHAPE_TYPE.RECTANGLE)
    shape = slide.shapes.add_shape(kind, Inches(e["x"]), Inches(e["y"]), Inches(e["w"]), Inches(e["h"]))
    shape.name = str(e["id"])
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(e.get("fill"), BG)
    stroke = e.get("stroke") or e.get("line_color")
    if stroke:
        shape.line.color.rgb = rgb(stroke)
        shape.line.width = Pt(float(e.get("line_width", 0.7)))
    else:
        shape.line.fill.background()
    return shape


def add_image(slide, e: dict):
    path = ASSET_ROOT / str(e["image_source"])
    if not path.is_file():
        raise RuntimeError(f"image unavailable during render: {e['image_source']}")
    shape = slide.shapes.add_picture(str(path), Inches(e["x"]), Inches(e["y"]), width=Inches(e["w"]), height=Inches(e["h"]))
    shape.name = str(e["id"])
    if str(e.get("crop", e.get("crop_mode", ""))).lower() in {"cover", "crop"}:
        with Image.open(path) as image:
            actual, target = image.width / image.height, e["w"] / e["h"]
        if actual > target:
            shape.crop_left = shape.crop_right = (1 - target / actual) / 2
        else:
            shape.crop_top = shape.crop_bottom = (1 - actual / target) / 2
    return shape


def add_line(slide, e: dict, positions: dict):
    source, target = e.get("from"), e.get("to")
    if isinstance(source, str) and isinstance(target, str):
        a, b = positions[source], positions[target]
        x1, y1 = a.left + a.width // 2, a.top + a.height // 2
        x2, y2 = b.left + b.width // 2, b.top + b.height // 2
    elif isinstance(source, list) and isinstance(target, list) and len(source) == len(target) == 2:
        x1, y1, x2, y2 = Inches(float(source[0])), Inches(float(source[1])), Inches(float(target[0])), Inches(float(target[1]))
    else:
        x1, y1 = Inches(e["x"]), Inches(e["y"])
        x2, y2 = Inches(e["x"] + e["w"]), Inches(e["y"] + e["h"])
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    shape.name = str(e["id"]); shape.line.color.rgb = rgb(e.get("color"), BLUE); shape.line.width = Pt(float(e.get("line_width", 1.2)))
    return shape


def decorate(slide, page_no: int) -> None:
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(*BG)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(.13), Inches(H))
    bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(*BLUE); bar.line.fill.background()
    add_text(slide, {"id": "brand_wordmark", "x": 11.6, "y": .12, "w": 1.1, "h": .18, "text": "COMKING", "font_size": 9, "font_weight": "bold", "color": "#1F497D", "alignment": "right"})
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(.65), Inches(6.91), Inches(12.05), Inches(.02))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(*LIGHT_BLUE); line.line.fill.background()
    add_text(slide, {"id": "footer", "x": .65, "y": 7.02, "w": 4.3, "h": .18, "text": "医院园区智慧配电改造方案", "font_size": 9, "color": "#5A6B78"})
    add_text(slide, {"id": "page_number", "x": 12.05, "y": 7.00, "w": .55, "h": .2, "text": str(page_no).zfill(2), "font_size": 10, "font_weight": "bold", "color": "#4F81BD", "alignment": "right"})


def draw_deck(scenes: dict[str, dict], deck: Path) -> None:
    prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    for letter in PAGES:
        slide = prs.slides.add_slide(blank); decorate(slide, PAGES[letter]); positions = {}
        elements = sorted(scenes[letter]["elements"], key=lambda e: float(e.get("z_order", e.get("z", 10))))
        for raw in elements:
            e = clamp(dict(raw)); typ = e["type"]
            if typ in {"text", "number", "caption"}: shape = add_text(slide, e)
            elif typ == "image": shape = add_image(slide, e)
            elif typ in {"rectangle", "rounded_rectangle", "circle", "callout"}: shape = add_shape(slide, e)
            elif typ in {"line", "arrow"}: shape = add_line(slide, e, positions)
            elif typ == "group": continue  # grouping is semantic; children must be emitted as elements.
            else: continue
            positions[str(e["id"])] = shape
    deck.parent.mkdir(parents=True, exist_ok=True); prs.save(deck)


def render(deck: Path, output: Path) -> list[Path]:
    output.mkdir(parents=True, exist_ok=True)
    subprocess.run(["powershell", "-ExecutionPolicy", "Bypass", "-File", str(ROOT / "scripts" / "render_pptx_windows.ps1"), "-InputPath", str(deck), "-OutputDir", str(output)], cwd=ROOT, check=True, timeout=420)
    image_dir = output / "slide" if (output / "slide").is_dir() else output
    pngs = sorted(image_dir.glob("*.png"), key=lambda p: int(re.search(r"(\d+)$", p.stem).group(1)) if re.search(r"(\d+)$", p.stem) else p.name)
    if len(pngs) != len(PAGES):
        raise RuntimeError(f"PowerPoint COM expected {len(PAGES)} PNGs, got {len(pngs)}")
    return pngs


def contact_sheet(pngs: list[Path], output: Path, labels: list[str]) -> None:
    images = [Image.open(path).convert("RGB") for path in pngs]
    width, gap, header = 400, 16, 34
    thumb_h = round(images[0].height * width / images[0].width)
    sheet = Image.new("RGB", (width * 3 + gap * 2, (thumb_h + header) * 2), "white")
    draw = ImageDraw.Draw(sheet)
    for index, image in enumerate(images):
        row, col = divmod(index, 3); x = col * (width + gap); y = row * (thumb_h + header)
        draw.text((x + 5, y + 5), labels[index], fill=TEXT)
        sheet.paste(image.resize((width, thumb_h)), (x, y + header))
    output.parent.mkdir(parents=True, exist_ok=True); sheet.save(output)


def find_slide_image(folder: Path, page_no: int) -> Path | None:
    if not folder.is_dir(): return None
    return next((path for path in folder.glob("*.png") if re.search(rf"{page_no}$", path.stem)), None)


def comparison(pure_pngs: list[Path], output: Path) -> None:
    rows = []
    for letter, pure in zip(PAGES, pure_pngs):
        page_no = PAGES[letter]
        rows.append((f"Hybrid V5 | {letter}", find_slide_image(HYBRID_PNG, page_no)))
        # Stabilization writes independently and currently supplies no
        # semantic A-F manifest. Slide filenames alone cannot prove identity,
        # so never substitute a merely same-numbered or legacy slide here.
        rows.append((f"Scene Graph Stabilization | {letter}", None))
        rows.append((f"Pure Art Director | {letter}", pure))
    width, gap, header, row_gap = 410, 12, 30, 16
    sample = Image.open(pure_pngs[0]); height = round(sample.height * width / sample.width)
    sheet = Image.new("RGB", (width * 3 + gap * 2, (height + header + row_gap) * len(PAGES)), "white")
    draw = ImageDraw.Draw(sheet)
    for index, (label, path) in enumerate(rows):
        row, col = divmod(index, 3); x = col * (width + gap); y = row * (height + header + row_gap)
        draw.text((x + 4, y + 5), label, fill=TEXT)
        if path and path.is_file():
            with Image.open(path) as image: sheet.paste(image.convert("RGB").resize((width, height)), (x, y + header))
        else:
            draw.rectangle((x, y + header, x + width, y + header + height), fill=(238, 241, 244), outline=(180, 190, 200))
            draw.text((x + width // 2 - 28, y + header + height // 2), "PENDING", fill=(80, 90, 100))
    output.parent.mkdir(parents=True, exist_ok=True); sheet.save(output)


def image_data(path: Path) -> str:
    return "data:image/png;base64," + base64.b64encode(path.read_bytes()).decode("ascii")


def revision_prompt(scene: dict, png: Path) -> list[dict]:
    # Revision is intentionally a single art-director self-review, after the
    # first pass is rendered. It receives no page-type or layout prescription.
    content = [
        {"type": "text", "text": "保持业务内容不变，你是一名 PPT Art Director。请判断当前页面视觉是否需要调整。只调整视觉布局，不要改变业务内容。只输出完整合法 JSON Scene Graph。"},
        {"type": "image_url", "image_url": {"url": image_data(png)}},
        {"type": "text", "text": "当前 Scene Graph：\n" + json.dumps(scene, ensure_ascii=False)},
    ]
    return [{"role": "user", "content": content}]


def drift(scene: dict, letter: str, briefs: dict[str, dict]) -> list[str]:
    """Record only obvious cross-page topic leakage; never rewrite it."""
    text = " ".join(str(e.get("text", "")) for e in scene.get("elements", []) if isinstance(e, dict))
    markers = {"A": ["预训练运维", "无人值守", "分阶段", "参数确认"], "B": ["无人值守", "分阶段", "参数确认"], "C": ["预训练运维", "分阶段", "参数确认"], "D": ["预训练运维", "无人值守", "参数确认"], "E": ["预训练运维", "无人值守", "分阶段", "参数确认"], "F": ["预训练运维", "无人值守", "分阶段"]}
    return [term for term in markers[letter] if term in text]


def changed_visual_fields(before: dict, after: dict) -> list[str]:
    old = {str(e.get("id")): e for e in before.get("elements", []) if isinstance(e, dict)}
    new = {str(e.get("id")): e for e in after.get("elements", []) if isinstance(e, dict)}
    changed = []
    for ident in sorted(set(old) | set(new)):
        if ident not in old or ident not in new:
            changed.append(f"{ident}: element added/removed")
            continue
        fields = [key for key in ("x", "y", "w", "h", "z_order", "z", "font_size", "color", "fill", "crop") if old[ident].get(key) != new[ident].get(key)]
        if fields: changed.append(f"{ident}: " + ", ".join(fields))
    return changed or ["no scene-graph visual changes returned"]


def load_or_generate_first(briefs: dict[str, dict], letters: list[str] | None = None) -> tuple[dict[str, dict], dict[str, dict]]:
    raw_dir, plan_dir = OUT / "raw", OUT / "plans"
    raw_dir.mkdir(parents=True, exist_ok=True); plan_dir.mkdir(parents=True, exist_ok=True)
    scenes, metadata = {}, {}
    for letter in letters or list(PAGES):
        brief = briefs[letter]
        raw_path, plan_path = raw_dir / f"page_{letter}_raw.json", plan_dir / f"page_{letter}_scene.json"
        if raw_path.is_file() and plan_path.is_file():
            scenes[letter] = json.loads(plan_path.read_text(encoding="utf-8")); metadata[letter] = {"status": "reused_existing_first_pass"}; continue
        raw, meta = ask(first_pass_prompt(brief), f"page_{letter}_first_pass")
        # Persist the untouched model response before parsing or validating it.
        raw_path.write_text(raw, encoding="utf-8")
        scene = as_json(raw)
        validation = validate(scene, {x["image_path"] for x in brief["image_assets"]})
        if validation["hard_errors"]:
            raise RuntimeError(f"page {letter} hard validation failed: {validation['hard_errors']}")
        if any(x["required"] and not any(e.get("type") == "image" and e.get("image_source") == x["image_path"] for e in scene["elements"]) for x in brief["image_assets"]):
            raise RuntimeError(f"page {letter} omitted a required real image")
        plan_path.write_text(json.dumps(scene, ensure_ascii=False, indent=2), encoding="utf-8")
        scenes[letter], metadata[letter] = scene, {**meta, "validation": validation}
    return scenes, metadata


def revise_all(scenes: dict[str, dict], briefs: dict[str, dict], first_pngs: list[Path], letters: list[str] | None = None) -> tuple[dict[str, dict], dict[str, dict]]:
    final_scenes, notes = {}, {}
    by_letter = dict(zip(PAGES, first_pngs))
    for letter in letters or list(PAGES):
        original = scenes[letter]
        raw_path, plan_path = OUT / "raw" / f"page_{letter}_revision_raw.json", OUT / "plans" / f"page_{letter}_scene_final.json"
        if raw_path.is_file() and plan_path.is_file():
            revised = json.loads(plan_path.read_text(encoding="utf-8")); meta = {"status": "reused_existing_revision"}
        else:
            raw, meta = ask(revision_prompt(original, by_letter[letter]), f"page_{letter}_visual_revision")
            raw_path.write_text(raw, encoding="utf-8")
            revised = as_json(raw)
            validation = validate(revised, {x["image_path"] for x in briefs[letter]["image_assets"]})
            if validation["hard_errors"]:
                raise RuntimeError(f"page {letter} revision hard validation failed: {validation['hard_errors']}")
            plan_path.write_text(json.dumps(revised, ensure_ascii=False, indent=2), encoding="utf-8")
        validation = validate(revised, {x["image_path"] for x in briefs[letter]["image_assets"]})
        final_scenes[letter] = revised
        notes[letter] = {"model": meta, "visual_changes": changed_visual_fields(original, revised), "validation": validation, "content_drift": drift(revised, letter, briefs)}
    return final_scenes, notes


def report(first_meta: dict, final_notes: dict, first_ppt: Path, final_ppt: Path) -> None:
    lines = ["# Pure Art Director experiment report", "", "## Tested pages", ""]
    for letter in PAGES:
        val = first_meta[letter].get("validation", {})
        scene = json.loads((OUT / "plans" / f"page_{letter}_scene.json").read_text(encoding="utf-8"))
        images = [e.get("image_source") for e in scene.get("elements", []) if e.get("type") == "image"]
        lines += [f"- {letter}: source page {PAGES[letter]}, {len(scene.get('elements', []))} elements; images: {', '.join(images) if images else 'none'}; first-pass bounds={val.get('preclamp_out_of_bounds', [])}; overlaps={val.get('overlap_pairs', [])}; small-font={val.get('small_font', [])}"]
    lines += ["", "## One visual self-revision", ""]
    for letter in PAGES:
        note = final_notes[letter]
        lines.append(f"- {letter}: {'; '.join(note['visual_changes'])}; CONTENT_DRIFT: {note['content_drift'] or 'none'}")
    lines += ["", "## Render outputs", "", f"- First pass: `{first_ppt.relative_to(ROOT)}`", f"- Final: `{final_ppt.relative_to(ROOT)}`", f"- COM rendering: six first-pass and six final PNGs were required by the runner.", "", "## Method boundary", "", "- Seed owned the page scene graphs. The renderer only executed JSON, added the common wordmark/footer, and applied minimal bounds clamping.", "- Unlike Stabilization, this route contains no semantic content lock, visual reference, template selector, page-type classifier, targeted layout rule, or automatic spacing repair."]
    (OUT / "pure_art_director_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def completed_scenes(suffix: str = "") -> dict[str, dict]:
    return {letter: json.loads((OUT / "plans" / f"page_{letter}_scene{suffix}.json").read_text(encoding="utf-8")) for letter in PAGES}


def main() -> None:
    parser = argparse.ArgumentParser(description="Run the isolated Pure Art Director experiment.")
    parser.add_argument("--first-page", choices=PAGES, help="Generate exactly one resumable first-pass page.")
    parser.add_argument("--revision-page", choices=PAGES, help="Generate exactly one rendered-PNG visual revision.")
    parser.add_argument("--render-first", action="store_true", help="Render completed first-pass scene files only.")
    parser.add_argument("--render-final", action="store_true", help="Render completed revision scene files only.")
    parser.add_argument("--report", action="store_true", help="Write the report from completed experiment artifacts only.")
    args = parser.parse_args()
    if sum(bool(x) for x in (args.first_page, args.revision_page, args.render_first, args.render_final, args.report)) > 1:
        parser.error("choose one resumable action at a time")
    OUT.mkdir(parents=True, exist_ok=True)
    briefs = source_briefs()
    (OUT / "plans" / "content_briefs.json").parent.mkdir(parents=True, exist_ok=True)
    (OUT / "plans" / "content_briefs.json").write_text(json.dumps(briefs, ensure_ascii=False, indent=2), encoding="utf-8")
    if args.first_page:
        load_or_generate_first(briefs, [args.first_page])
        return
    if args.render_first:
        scenes = completed_scenes()
        first_dir = OUT / "first_pass"; first_ppt = first_dir / "pure_art_director_first_pass.pptx"
        draw_deck(scenes, first_ppt)
        first_pngs = render(first_ppt, first_dir / "png")
        contact_sheet(first_pngs, first_dir / "contact_sheet.png", [f"Pure Art Director | {x}" for x in PAGES])
        comparison(first_pngs, OUT / "comparison.png")
        return
    if args.revision_page:
        scenes = completed_scenes()
        first_pngs = sorted((OUT / "first_pass" / "png" / "slide").glob("*.png"), key=lambda p: int(re.search(r"(\d+)$", p.stem).group(1)))
        if len(first_pngs) != len(PAGES):
            raise RuntimeError("render the complete first pass before a visual revision")
        revise_all(scenes, briefs, first_pngs, [args.revision_page])
        return
    if args.render_final:
        final_scenes = completed_scenes("_final")
        final_dir = OUT / "final"; final_ppt = final_dir / "pure_art_director_final.pptx"
        draw_deck(final_scenes, final_ppt)
        final_pngs = render(final_ppt, final_dir / "png")
        contact_sheet(final_pngs, final_dir / "contact_sheet.png", [f"Pure Art Director final | {x}" for x in PAGES])
        comparison(final_pngs, OUT / "comparison_final.png")
        return
    if args.report:
        scenes = completed_scenes()
        first_meta = {letter: {"validation": validate(scenes[letter], {x["image_path"] for x in briefs[letter]["image_assets"]})} for letter in PAGES}
        first_pngs = sorted((OUT / "first_pass" / "png" / "slide").glob("*.png"), key=lambda p: int(re.search(r"(\d+)$", p.stem).group(1)))
        _, final_notes = revise_all(scenes, briefs, first_pngs)
        report(first_meta, final_notes, OUT / "first_pass" / "pure_art_director_first_pass.pptx", OUT / "final" / "pure_art_director_final.pptx")
        return
    scenes, first_meta = load_or_generate_first(briefs)
    first_dir = OUT / "first_pass"; first_ppt = first_dir / "pure_art_director_first_pass.pptx"
    draw_deck(scenes, first_ppt)
    first_pngs = render(first_ppt, first_dir / "png")
    contact_sheet(first_pngs, first_dir / "contact_sheet.png", [f"Pure Art Director | {x}" for x in PAGES])
    comparison(first_pngs, OUT / "comparison.png")
    final_scenes, final_notes = revise_all(scenes, briefs, first_pngs)
    final_dir = OUT / "final"; final_ppt = final_dir / "pure_art_director_final.pptx"
    draw_deck(final_scenes, final_ppt)
    final_pngs = render(final_ppt, final_dir / "png")
    contact_sheet(final_pngs, final_dir / "contact_sheet.png", [f"Pure Art Director final | {x}" for x in PAGES])
    comparison(final_pngs, OUT / "comparison_final.png")
    report(first_meta, final_notes, first_ppt, final_ppt)
    print(json.dumps({"pages": list(PAGES), "first_pass_pptx": str(first_ppt), "final_pptx": str(final_ppt)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
