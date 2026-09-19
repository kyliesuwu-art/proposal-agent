"""Resume the Pure Art Director experiment as a single 15-page deck.

Only E and G--O can call Ark. A--C/F are reused, D is locally corrected.
Each new response is written before JSON parsing/validation.
"""
from __future__ import annotations

import importlib.util
import json
import re
import subprocess
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent.parent
OLD = ROOT / "outputs" / "ppt_pure_art_director"
OUT = OLD / "full15"
SOURCE = ROOT / "outputs" / "ppt_style_v4" / "slides.final.md"
ASSETS = ROOT / "outputs" / "ppt_style_v4"
LETTERS = tuple("ABCDEFGHIJKLMNO")
# Existing A--F are intentionally retained in their already-established
# semantic order; new pages provide the missing narrative context and close.
SOURCE_PAGES = dict(zip(LETTERS, (9, 10, 5, 12, 13, 14, 1, 2, 3, 4, 6, 7, 8, 11, 15)))
IMAGE_OVERRIDES = {"E": ["assets/image-005.png"], "G": ["assets/image-001.jpg"]}


def pad_module():
    spec = importlib.util.spec_from_file_location("pure_art", ROOT / "scripts" / "run_ppt_pure_art_director.py")
    module = importlib.util.module_from_spec(spec)
    assert spec.loader
    spec.loader.exec_module(module)
    return module


PAD = pad_module()


def pieces() -> list[str]:
    return [x.strip() for x in SOURCE.read_text(encoding="utf-8").replace("\r\n", "\n").split("\n---\n") if x.strip()]


def brief(letter: str, piece: str) -> dict:
    title = next((x.lstrip("#").strip() for x in piece.splitlines() if x.startswith("#")), letter)
    bullets = [re.sub(r"^\s*[*-]\s*", "", x).strip() for x in piece.splitlines() if re.match(r"^\s*[*-]\s+", x)]
    images = re.findall(r"!\[[^]]*\]\((assets/[^)\s]+)\)", piece)
    images = IMAGE_OVERRIDES.get(letter, images)
    return {"source_page": SOURCE_PAGES[letter], "title": title, "primary_message": bullets[0] if bullets else title,
            "facts_and_bullets": bullets[:5], "image_assets": images}


def visual_intent(source_page: int) -> dict:
    plan = json.loads((ROOT / "outputs" / "ppt_style_v4" / "visual_plan.json").read_text(encoding="utf-8"))
    return next((x for x in plan["plan"] if x["slide"] == source_page), {})


def prompt(letter: str, info: dict) -> list[dict]:
    required = info["image_assets"]
    system = """You are the art director of one formal 16:9 COMKING hospital-energy presentation page. Return only one JSON object with background and elements. Canvas: 13.333 by 7.5 inches. Each element has id,type,x,y,w,h; use only text,image,rectangle,rounded_rectangle,circle,line,arrow. Exact schema is mandatory: text elements use text, font_size, font_weight, alignment; image elements use image_source (never src) and crop:'cover'; lines/arrows require positive w/h or valid from/to coordinates. Do not use aliases such as content, fontSize, textAlign, strokeWidth, or src. Text must be concise, no Markdown, URLs, citations or invented facts. Use COMKING blue/light blue, clear hierarchy, generous but purposeful whitespace, and body type >=16pt. Images must be a main composition element (not a corner thumbnail), must preserve aspect ratio with crop:'cover', and may only use listed assets. All geometry, including arrows, must remain within the canvas. Do not emit page number, logo or footer; deterministic renderer adds them."""
    payload = {"page_id": letter, "content": {k: info[k] for k in ("title", "primary_message", "facts_and_bullets")},
               "approved_images": required, "required_images": required,
               "visual_intent": visual_intent(info["source_page"])}
    return [{"role": "system", "content": system}, {"role": "user", "content": json.dumps(payload, ensure_ascii=False)}]


def validate(scene: dict, info: dict) -> dict:
    result = PAD.validate(scene, set(info["image_assets"]))
    used = {x.get("image_source") for x in scene.get("elements", []) if x.get("type") == "image"}
    missing = [x for x in info["image_assets"] if x not in used]
    if missing:
        result["hard_errors"].append("required_image_missing:" + ",".join(missing))
    return result


def local_d(scene: dict) -> dict:
    fixed = json.loads(json.dumps(scene))
    for e in fixed["elements"]:
        if e.get("id") == "arrow_1to2": e["from"], e["to"] = [4.36, 3.5], [4.70, 3.5]
        if e.get("id") == "arrow_2to3": e["from"], e["to"] = [8.56, 3.5], [8.90, 3.5]
    return fixed


def obtain(letter: str, info: dict) -> tuple[dict, str]:
    plan = OUT / "plans" / f"page_{letter}_scene.json"
    raw = OUT / "raw" / f"page_{letter}_raw.json"
    if plan.is_file() and raw.is_file():
        return json.loads(plan.read_text(encoding="utf-8")), "checkpoint_reused"
    if letter in "ABCF":
        scene = json.loads((OLD / "plans" / f"page_{letter}_scene_final.json").read_text(encoding="utf-8")); status = "reused_existing"
    elif letter == "D":
        scene = local_d(json.loads((OLD / "plans" / "page_D_scene_final.json").read_text(encoding="utf-8"))); status = "locally_corrected"
    else:
        text, _ = PAD.ask(prompt(letter, info), f"full15_{letter}")
        raw.parent.mkdir(parents=True, exist_ok=True); raw.write_text(text, encoding="utf-8")
        scene, status = PAD.as_json(text), "ark_generated"
    check = validate(scene, info)
    if check["hard_errors"]:
        raise RuntimeError(f"{letter} validation failed: {check['hard_errors']}")
    plan.parent.mkdir(parents=True, exist_ok=True); plan.write_text(json.dumps(scene, ensure_ascii=False, indent=2), encoding="utf-8")
    return scene, status


def draw(scenes: dict[str, dict], target: Path) -> None:
    prs = Presentation(); prs.slide_width = Inches(PAD.W); prs.slide_height = Inches(PAD.H); blank = prs.slide_layouts[6]
    for number, letter in enumerate(LETTERS, 1):
        slide = prs.slides.add_slide(blank); PAD.decorate(slide, number); positions = {}
        for raw in sorted(scenes[letter]["elements"], key=lambda e: float(e.get("z_order", e.get("z", 10)))):
            e = PAD.clamp(dict(raw)); typ = e["type"]
            if typ in {"text", "number", "caption"}: shape = PAD.add_text(slide, e)
            elif typ == "image": shape = PAD.add_image(slide, e)
            elif typ in {"rectangle", "rounded_rectangle", "circle", "callout"}: shape = PAD.add_shape(slide, e)
            elif typ in {"line", "arrow"}: shape = PAD.add_line(slide, e, positions)
            else: continue
            positions[str(e["id"])] = shape
    target.parent.mkdir(parents=True, exist_ok=True); prs.save(target)


def render(deck: Path) -> list[Path]:
    out = OUT / "rendered" / "png"; out.mkdir(parents=True, exist_ok=True)
    subprocess.run(["powershell", "-ExecutionPolicy", "Bypass", "-File", str(ROOT / "scripts" / "render_pptx_windows.ps1"), "-InputPath", str(deck), "-OutputDir", str(out)], cwd=ROOT, check=True, timeout=420)
    images = sorted((out / "slide").glob("*.png"), key=lambda p: int(re.search(r"(\d+)$", p.stem).group(1)))
    if len(images) != 15: raise RuntimeError(f"PowerPoint COM rendered {len(images)}/15 pages")
    return images


def contact(images: list[Path]) -> None:
    sample = Image.open(images[0]).convert("RGB"); width, gap, head = 340, 12, 28; height = round(sample.height * width / sample.width)
    sheet = Image.new("RGB", (width * 3 + gap * 2, (height + head) * 5), "white"); pen = ImageDraw.Draw(sheet)
    for i, image in enumerate(images):
        row, col = divmod(i, 3); x, y = col * (width + gap), row * (height + head); pen.text((x + 4, y + 4), f"{LETTERS[i]} | {i+1:02}", fill=(24, 50, 71))
        sheet.paste(Image.open(image).convert("RGB").resize((width, height)), (x, y + head))
    sheet.save(OUT / "rendered" / "contact_sheet.png")


def main() -> None:
    all_pieces = pieces(); infos = {x: brief(x, all_pieces[SOURCE_PAGES[x]-1]) for x in LETTERS}
    (OUT / "content_briefs.json").parent.mkdir(parents=True, exist_ok=True); (OUT / "content_briefs.json").write_text(json.dumps(infos, ensure_ascii=False, indent=2), encoding="utf-8")
    scenes, states = {}, {}
    for letter in LETTERS: scenes[letter], states[letter] = obtain(letter, infos[letter])
    deck = OUT / "pure_art_director_full15.pptx"; draw(scenes, deck); images = render(deck); contact(images)
    audit = {x: validate(scenes[x], infos[x]) for x in LETTERS}
    (OUT / "automatic_audit.json").write_text(json.dumps({"pages": states, "validation": audit, "image_driven_pages": [x for x in LETTERS if infos[x]["image_assets"]]}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"deck": str(deck), "pages": states}, ensure_ascii=False))


if __name__ == "__main__": main()
